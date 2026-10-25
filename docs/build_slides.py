"""
build_slides.py — 產生「實驗室本地 LLM 使用教學」投影片（docs/實驗室本地LLM使用教學.pptx）

修改內容後重新產生（在 repo 根目錄；借用 Discord bot 的 Python 環境，裡面有 python-pptx，
先跑過 discord\\client\\deploy.ps1）：
    discord\\client\\.venv\\Scripts\\python.exe docs\\build_slides.py

實際畫面的截圖放在 docs/img/，換截圖時直接覆蓋同名檔案再重新產生。
"""

import os

from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

OUT = os.path.join(os.path.dirname(os.path.abspath(__file__)), "實驗室本地LLM使用教學.pptx")
IMG = os.path.join(os.path.dirname(os.path.abspath(__file__)), "img")

# --- 配色（沿用舊版投影片）與字型 -----------------------------------------
NAVY   = RGBColor(0x0F, 0x1B, 0x2D)   # 深色頁背景
NAVY2  = RGBColor(0x1C, 0x2B, 0x44)   # 深色頁上的卡片
INK    = RGBColor(0x1B, 0x24, 0x33)   # 標題文字
BODY   = RGBColor(0x4A, 0x55, 0x68)   # 內文
MUTED  = RGBColor(0x8B, 0x9B, 0xB4)   # 註解 / 頁尾
BLUE   = RGBColor(0x25, 0x63, 0xC9)   # 主要強調色
ORANGE = RGBColor(0xF0, 0xA0, 0x4B)   # 第二強調色
BG     = RGBColor(0xF6, 0xF5, 0xF1)   # 內容頁背景
CARD   = RGBColor(0xFD, 0xFC, 0xF9)   # 卡片
BORDER = RGBColor(0xE2, 0xDF, 0xD8)
TINT   = RGBColor(0xE6, 0xED, 0xF8)   # 表頭 / 淡藍底
CODEBG = RGBColor(0x1B, 0x24, 0x33)
CODETX = RGBColor(0xE6, 0xE9, 0xEF)
WHITE  = RGBColor(0xFF, 0xFF, 0xFF)
RED    = RGBColor(0xC2, 0x41, 0x0C)

FONT = "微軟正黑體"
MONO = "Consolas"

W, H = Inches(13.333), Inches(7.5)
MX = Inches(0.6)                      # 左右邊界
CW = W - MX * 2                       # 內容寬度
DECK = "實驗室本地 LLM 使用教學"

prs = Presentation()
prs.slide_width, prs.slide_height = W, H
BLANK = prs.slide_layouts[6]
_page = [0]


# =============================================================================
#  基本工具
# =============================================================================
def _font(run, size, color, bold=False, mono=False):
    f = run.font
    f.size = Pt(size)
    f.bold = bold
    f.color.rgb = color
    f.name = MONO if mono else FONT
    rPr = run._r.get_or_add_rPr()
    for old in rPr.findall(qn("a:ea")):
        rPr.remove(old)
    ea = etree.SubElement(rPr, qn("a:ea"))
    ea.set("typeface", FONT)           # 中文一律微軟正黑體（程式碼區的中文註解也是）
    latin = rPr.find(qn("a:latin"))
    if latin is not None:
        latin.addnext(ea)


def text(slide, x, y, w, h, paras, size=16, color=BODY, bold=False, mono=False,
         align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, space=6, line=None):
    """paras：字串、或段落清單；每段是字串，或 [(文字, {size,color,bold,mono}), ...]"""
    tb = slide.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    if isinstance(paras, str):
        paras = [paras]
    for i, para in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.space_after = Pt(space)
        if line:
            p.line_spacing = line
        runs = para if isinstance(para, list) else [(para, {})]
        for t, o in runs:
            r = p.add_run()
            r.text = t
            _font(r, o.get("size", size), o.get("color", color), o.get("bold", bold), o.get("mono", mono))
    return tb


def box(slide, x, y, w, h, fill=CARD, line=BORDER, radius=0.06, shape=MSO_SHAPE.ROUNDED_RECTANGLE):
    s = slide.shapes.add_shape(shape, x, y, w, h)
    s.fill.solid()
    s.fill.fore_color.rgb = fill
    if line is None:
        s.line.fill.background()
    else:
        s.line.color.rgb = line
        s.line.width = Pt(0.75)
    s.shadow.inherit = False
    if shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        s.adjustments[0] = radius
    s.text_frame.text = ""
    return s


def badge(slide, x, y, d, label, fill=BLUE, color=WHITE, size=14):
    c = box(slide, x, y, d, d, fill=fill, line=None, shape=MSO_SHAPE.OVAL)
    tf = c.text_frame
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = label
    _font(r, size, color, bold=True)
    return c


def arrow(slide, x1, y1, x2, y2, color=MUTED, width=1.5, dash=False):
    # 座標一定要是整數 EMU：連接線不會自動取整，小數會讓 PowerPoint 打不開檔案
    c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, int(x1), int(y1), int(x2), int(y2))
    c.line.color.rgb = color
    c.line.width = Pt(width)
    ln = c.line._get_or_add_ln()
    if dash:
        d = etree.SubElement(ln, qn("a:prstDash"))
        d.set("val", "dash")
    tail = etree.SubElement(ln, qn("a:tailEnd"))
    tail.set("type", "triangle")
    return c


def screenshot(slide, x, y, w, name):
    """docs/img/ 裡的截圖，寬度固定、高度照比例，加細框"""
    pic = slide.shapes.add_picture(os.path.join(IMG, name), int(x), int(y), width=int(w))
    pic.line.color.rgb = BORDER
    pic.line.width = Pt(0.75)
    return pic


def highlight(slide, pic, fx, fy, fw, fh, label):
    """在截圖上框出一塊（fx…fh 是佔圖片寬高的比例），右上角放編號"""
    x, y = pic.left + int(pic.width * fx), pic.top + int(pic.height * fy)
    w, h = int(pic.width * fw), int(pic.height * fh)
    s = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h)
    s.fill.background()
    s.line.color.rgb = ORANGE
    s.line.width = Pt(2.25)
    s.shadow.inherit = False
    s.adjustments[0] = 0.08
    d = Inches(0.34)
    badge(slide, x + w - d // 2, y - d // 2, d, label, fill=ORANGE, size=12)


def code(slide, x, y, w, h, lines, size=13):
    """深色程式碼區塊。# 開頭 = 註解（灰）；> 開頭 = 重點（橘，> 不會顯示）"""
    box(slide, x, y, w, h, fill=CODEBG, line=None, radius=0.04)
    paras = []
    for ln in lines:
        if ln.startswith("#"):
            paras.append([(ln, {"color": MUTED})])
        elif ln.startswith(">"):
            paras.append([(ln[1:].lstrip(), {"color": ORANGE, "bold": True})])
        else:
            paras.append([(ln, {"color": CODETX})])
    text(slide, x + Inches(0.3), y + Inches(0.25), w - Inches(0.6), h - Inches(0.5),
         paras, size=size, mono=True, space=4)


def table(slide, x, y, w, rows, colw, size=13, rowh=0.42, header=True, mono_cols=()):
    nrows, ncols = len(rows), len(rows[0])
    gt = slide.shapes.add_table(nrows, ncols, x, y, w, Inches(rowh * nrows))
    tbl = gt.table
    tblPr = tbl._tbl.tblPr
    style = tblPr.find(qn("a:tableStyleId"))
    if style is not None:
        style.text = "{5940675A-B579-460E-94D1-54222C63F5DA}"   # 無樣式、表格格線
    tblPr.set("firstRow", "0")
    tblPr.set("bandRow", "0")
    total = sum(colw)
    for i, cw in enumerate(colw):
        tbl.columns[i].width = int(w * cw / total)
    for r in range(nrows):
        tbl.rows[r].height = Inches(rowh)
        for c in range(ncols):
            cell = tbl.cell(r, c)
            is_head = header and r == 0
            cell.fill.solid()
            cell.fill.fore_color.rgb = TINT if is_head else CARD
            cell.margin_left = cell.margin_right = Inches(0.12)
            cell.margin_top = cell.margin_bottom = Inches(0.05)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            tf = cell.text_frame
            tf.word_wrap = True
            p = tf.paragraphs[0]
            r_ = p.add_run()
            r_.text = rows[r][c]
            _font(r_, size, INK if is_head else BODY, bold=is_head,
                  mono=(c in mono_cols and not is_head))
            # 框線（a:lnL…）在 schema 裡必須排在填色（a:solidFill）之前，所以插在最前面
            tcPr = cell._tc.get_or_add_tcPr()
            for i, tag in enumerate(("a:lnL", "a:lnR", "a:lnT", "a:lnB")):
                ln = etree.Element(qn(tag))
                ln.set("w", "9525")
                sf = etree.SubElement(ln, qn("a:solidFill"))
                clr = etree.SubElement(sf, qn("a:srgbClr"))
                clr.set("val", "E2DFD8")
                tcPr.insert(i, ln)
    return gt


# =============================================================================
#  頁面骨架
# =============================================================================
def new_slide(dark=False):
    s = prs.slides.add_slide(BLANK)
    s.background.fill.solid()
    s.background.fill.fore_color.rgb = NAVY if dark else BG
    _page[0] += 1
    return s


def header(s, kicker, title, dark=False):
    text(s, MX, Inches(0.45), CW, Inches(0.3), kicker, size=12, color=ORANGE if dark else BLUE, bold=True)
    text(s, MX, Inches(0.78), CW, Inches(0.7), title, size=30, color=WHITE if dark else INK, bold=True)


def footer(s, dark=False):
    text(s, MX, H - Inches(0.5), Inches(6), Inches(0.25), f"{DECK} · {_page[0]:02d}",
         size=10, color=MUTED)


def notes(s, t):
    s.notes_slide.notes_text_frame.text = t


def content_slide(kicker, title, note):
    s = new_slide()
    header(s, kicker, title)
    footer(s)
    notes(s, note)
    return s


TOP = Inches(1.75)   # 內容區起點


# =============================================================================
#  1. 封面
# =============================================================================
s = new_slide(dark=True)
text(s, MX, Inches(2.0), CW, Inches(0.35), "實驗室本地 LLM · 使用教學", size=14, color=ORANGE, bold=True)
text(s, MX, Inches(2.45), CW, Inches(1.0), "自架 vLLM × opencode × Discord", size=44, color=WHITE, bold=True)
text(s, MX, Inches(3.55), CW, Inches(0.5),
     "一台 Windows 筆電跑模型，同學用 opencode 寫程式、用 Discord 產生 PPT / Word / Excel",
     size=18, color=RGBColor(0xBF, 0xC8, 0xD6))
x = MX
for chip in ("vLLM :8000", "opencode", "Discord bot", "Windows + WSL2", "學校 VPN / Funnel"):
    wch = Inches(0.35 + len(chip) * 0.13)
    c = box(s, x, Inches(4.45), wch, Inches(0.42), fill=NAVY2, line=None, radius=0.5)
    text(s, x, Inches(4.45), wch, Inches(0.42), chip, size=12, color=CODETX,
         align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, space=0)
    x += wch + Inches(0.15)
text(s, MX, H - Inches(0.75), CW, Inches(0.3),
     "GitHub：CYH4157/local-llm · 完整步驟：完整步驟.md", size=11, color=MUTED)
notes(s, "這份教學說明實驗室怎麼用一台 Windows 筆電架設語言模型（vLLM），"
         "讓同學透過 opencode 寫程式、寫作業，並透過 Discord bot 產生簡報、Word、Excel 草稿。"
         "模型、程式碼與對話都留在實驗室，不經過雲端 API。")

# =============================================================================
#  2. 目錄
# =============================================================================
s = content_slide("目錄", "今天的內容",
                  "前兩段是系統介紹與一次性的主機安裝，第三、四段是每天會用到的 opencode 與 Discord bot，"
                  "第五段是不在實驗室時怎麼連，最後是管理者需要知道的金鑰、換模型與排錯。")
items = [("01", "認識系統", "三個元件、整體架構、模型實測"),
         ("02", "筆電主機安裝", "雙擊 setup-host.cmd 一次完成"),
         ("03", "同學使用 opencode", "安裝 client、基本操作、提問技巧"),
         ("04", "Discord bot", "設定、指令、上線前檢查"),
         ("05", "不在實驗室時", "在家用學校 VPN、校外用 Funnel"),
         ("06", "管理與排錯", "金鑰、換模型、排錯、解除安裝")]
cw, ch, gap = (CW - Inches(0.5)) / 3, Inches(2.1), Inches(0.25)
for i, (n, t, d) in enumerate(items):
    x = MX + (cw + gap) * (i % 3)
    y = TOP + (ch + gap) * (i // 3)
    box(s, x, y, cw, ch)
    text(s, x + Inches(0.35), y + Inches(0.35), cw - Inches(0.7), Inches(0.5), n, size=24, color=BLUE, bold=True)
    text(s, x + Inches(0.35), y + Inches(0.95), cw - Inches(0.7), Inches(0.45), t, size=20, color=INK, bold=True)
    text(s, x + Inches(0.35), y + Inches(1.45), cw - Inches(0.7), Inches(0.5), d, size=14, color=BODY)

# =============================================================================
#  3. 三個元件
# =============================================================================
s = content_slide("01 認識系統", "三個元件各司其職",
                  "vLLM 是唯一跑模型的地方，在筆電的 WSL 裡。opencode 和 Discord bot 都只是呼叫它的程式，"
                  "所以兩者共用同一個模型和同一張顯卡。同學的電腦不需要顯卡。")
comps = [("🧠", "vLLM", "推論伺服器：把模型載入 GPU，提供 OpenAI 相容 API，支援 tool calling。",
          "筆電 WSL · 服務 vllm · 埠 8000"),
         ("💻", "opencode", "AI 助手（桌面版 / 終端機版）：讀寫專案檔案、執行指令、寫程式、查資料。",
          "每台電腦都可以裝"),
         ("💬", "Discord bot", "在 Discord 問答、上網查詢，並產生 PPT / Word / Excel 草稿檔案。",
          "跑在筆電 · 自動跟著 vLLM")]
cw = (CW - Inches(0.5)) / 3
for i, (ic, t, d, foot) in enumerate(comps):
    x = MX + (cw + Inches(0.25)) * i
    box(s, x, TOP, cw, Inches(3.4))
    badge(s, x + Inches(0.35), TOP + Inches(0.35), Inches(0.7), ic, fill=TINT, size=20)
    text(s, x + Inches(0.35), TOP + Inches(1.25), cw - Inches(0.7), Inches(0.5), t, size=22, color=INK, bold=True)
    text(s, x + Inches(0.35), TOP + Inches(1.8), cw - Inches(0.7), Inches(1.0), d, size=15, color=BODY, line=1.15)
    text(s, x + Inches(0.35), TOP + Inches(2.85), cw - Inches(0.7), Inches(0.3), foot, size=12, color=BLUE, bold=True)
text(s, MX, TOP + Inches(3.8), CW, Inches(0.5),
     [[("模型、程式碼與對話都留在實驗室，", {"bold": True, "color": INK}),
       ("不經過任何雲端 API。", {"bold": True, "color": BLUE})]], size=17)

# =============================================================================
#  4. 整體架構
# =============================================================================
s = content_slide("01 認識系統", "整體架構",
                  "左邊是筆電：WSL 裡跑 vLLM，筆電本機的 opencode 和 Discord bot 用 127.0.0.1 連它。"
                  "右邊是三種使用者：在實驗室的同學直接用區網 IP；在家的同學先連學校 VPN，之後一樣用區網 IP；"
                  "沒有學校帳號的校外合作者，才用 Tailscale Funnel 的公開 HTTPS 網址。")
# 筆電
lx, ly, lw, lh = MX, TOP, Inches(5.4), Inches(4.6)
box(s, lx, ly, lw, lh, fill=TINT, line=None)
text(s, lx + Inches(0.3), ly + Inches(0.2), lw, Inches(0.4), "Windows 筆電（主機）", size=16, color=INK, bold=True)
box(s, lx + Inches(0.3), ly + Inches(0.75), lw - Inches(0.6), Inches(1.3), fill=CODEBG, line=None)
text(s, lx + Inches(0.55), ly + Inches(0.9), lw - Inches(1.1), Inches(1.1),
     [[("WSL2 Ubuntu", {"color": MUTED, "size": 12})],
      [("vLLM :8000", {"color": WHITE, "bold": True, "size": 20})],
      [("需 API 金鑰 · Qwen3-4B · 32k", {"color": CODETX, "size": 12})]], space=2)
vy = ly + Inches(2.05)
for i, (t, d) in enumerate([("opencode（本機）", "127.0.0.1:8000"), ("Discord bot", "127.0.0.1:8000 · 自動讀金鑰")]):
    bx = lx + Inches(0.3) + (lw - Inches(0.6)) / 2 * i + (Inches(0.1) if i else 0)
    bw = (lw - Inches(0.6)) / 2 - Inches(0.1)
    arrow(s, bx + bw / 2, ly + Inches(2.85), bx + bw / 2, vy + Inches(0.02), color=BLUE)
    box(s, bx, ly + Inches(2.85), bw, Inches(1.05))
    text(s, bx + Inches(0.2), ly + Inches(3.0), bw - Inches(0.4), Inches(0.8),
         [[(t, {"bold": True, "color": INK, "size": 15})], [(d, {"size": 11, "color": BODY})]], space=2)
text(s, lx + Inches(0.3), ly + Inches(4.05), lw - Inches(0.6), Inches(0.4),
     "防火牆 8000：只開放實驗室網段 + 學校 VPN 網段", size=12, color=BODY)
# 使用者
users = [("🏫", "實驗室同學（在校內）", "http://筆電IP:8000/v1", BLUE, False),
         ("🏠", "實驗室同學（在家）", "先連學校 VPN → 同上", BLUE, False),
         ("🌐", "校外合作者（無學校帳號）", "https://筆電.xxx.ts.net/v1（Funnel）", ORANGE, True)]
ux, uw, uh = MX + Inches(7.2), Inches(4.9), Inches(1.25)
for i, (ic, t, d, col, dash) in enumerate(users):
    uy = TOP + Inches(0.15) + (uh + Inches(0.3)) * i
    box(s, ux, uy, uw, uh)
    badge(s, ux + Inches(0.25), uy + Inches(0.3), Inches(0.65), ic, fill=TINT, size=18)
    text(s, ux + Inches(1.1), uy + Inches(0.22), uw - Inches(1.3), Inches(0.9),
         [[(t, {"bold": True, "color": INK, "size": 15})], [(d, {"size": 12, "color": col, "mono": True})]], space=3)
    arrow(s, ux - Inches(0.05), uy + uh / 2, lx + lw + Inches(0.05), ly + Inches(1.4), color=col, dash=dash)
text(s, ux, TOP + Inches(4.7), uw, Inches(0.3), "opencode 只需要「網址 + API 金鑰」", size=12, color=MUTED)

# =============================================================================
#  5. 模型實測
# =============================================================================
s = content_slide("01 認識系統 · 實驗室筆電 RTX 5070 Ti 12GB 實測", "預設模型：Qwen3-4B，32k 上下文",
                  "我們實測過 4B 和 8B。8B 沒有比較好：寫程式驗證錯了一半，長文件慢一倍，能同時服務的人也比較少。"
                  "所以預設用 4B（2507 新版），上下文開到 32k，因為 opencode 每次請求光系統提示就要約 1 萬 token。")
stats = [("27.5", "tok/s", "單人生成速度"), ("121", "tok/s", "五人同時總吞吐"), ("32k", "tokens", "單一對話上下文")]
sw = (CW - Inches(0.5)) / 3
for i, (n, u, d) in enumerate(stats):
    x = MX + (sw + Inches(0.25)) * i
    box(s, x, TOP, sw, Inches(1.75))
    text(s, x + Inches(0.35), TOP + Inches(0.2), sw, Inches(0.9),
         [[(n, {"size": 48, "bold": True, "color": BLUE}), ("  " + u, {"size": 16, "color": BODY})]], space=0)
    text(s, x + Inches(0.35), TOP + Inches(1.15), sw, Inches(0.4), d, size=14, color=INK, bold=True)
table(s, MX, TOP + Inches(2.1), CW, [
    ["測試", "4B · 16k", "4B · 32k（預設）", "8B-AWQ · 32k"],
    ["長文件 2.1 萬 token", "❌ 超過上限", "✅ 9.6 秒", "✅ 17.2 秒"],
    ["五人同時（總吞吐 / 最慢等待）", "61 tok/s / 6.7 秒", "121 tok/s / 3.4 秒", "98 tok/s / 4.1 秒"],
    ["寫程式（執行驗證 6 個案例）", "—", "✅ 6 / 6", "❌ 3 / 6"],
], [3.2, 2.2, 2.4, 2.2], size=14, rowh=0.5)
text(s, MX, TOP + Inches(4.35), CW, Inches(0.35),
     "8GB 顯卡請用 small 組合（4B · 16k）；換模型方式見第 06 段。", size=13, color=MUTED)

# =============================================================================
#  6. 開始前的準備
# =============================================================================
s = content_slide("02 筆電主機安裝", "開始前的準備",
                  "主機只需要一台有 NVIDIA 顯卡的 Windows 11 筆電。最容易忽略的是網路：筆電要接實驗室網路，"
                  "不能接手機熱點，而且最好請網管固定 IP，不然 IP 一變，所有人都要重設。")
reqs = [("🪟", "Windows 11 22H2 以上", "WSL mirrored 網路需要"),
        ("🎮", "NVIDIA 顯卡 8GB 以上", "裝最新版驅動；不要在 WSL 裡另裝"),
        ("💾", "硬碟約 30GB", "WSL、vLLM 約 8GB、模型 4~6GB"),
        ("🔐", "系統管理員權限", "只有安裝時需要"),
        ("🌐", "接實驗室網路、固定 IP", "不要用手機熱點（172.20.10.x）"),
        ("🔌", "插電使用", "腳本會設定插電時不睡眠")]
cw, ch = (CW - Inches(0.5)) / 3, Inches(2.05)
for i, (ic, t, d) in enumerate(reqs):
    x = MX + (cw + Inches(0.25)) * (i % 3)
    y = TOP + (ch + Inches(0.25)) * (i // 3)
    box(s, x, y, cw, ch)
    badge(s, x + Inches(0.3), y + Inches(0.3), Inches(0.6), ic, fill=TINT, size=16)
    text(s, x + Inches(0.3), y + Inches(1.05), cw - Inches(0.6), Inches(0.45), t, size=17, color=INK, bold=True)
    text(s, x + Inches(0.3), y + Inches(1.5), cw - Inches(0.6), Inches(0.45), d, size=13, color=BODY)

# =============================================================================
#  7. 一鍵安裝
# =============================================================================
s = content_slide("02 筆電主機安裝 · 全程約 20~40 分鐘，下載約 15GB", "雙擊 setup-host.cmd 完成安裝",
                  "整個安裝只要雙擊一個檔案。第一次如果還沒有 WSL，腳本會先安裝 Ubuntu 並要求重開機，"
                  "重開機後從開始選單開 Ubuntu 設定帳號密碼，再雙擊一次就會繼續。結束時會印出要給同學的網址和金鑰。")
steps = ["把專案資料夾放到筆電（git clone 或複製）",
         "雙擊 server\\wsl2\\setup-host.cmd，同意系統管理員權限",
         "第一次沒有 WSL：自動安裝 → 重開機 → 開「Ubuntu」設定帳密 → 再雙擊一次",
         "出現 [sudo] password 時輸入 Ubuntu 密碼（打字不會顯示）",
         "結束時記下印出的「網址」與「API 金鑰」"]
for i, st in enumerate(steps):
    y = TOP + Inches(0.92) * i
    badge(s, MX, y + Inches(0.05), Inches(0.5), str(i + 1), size=14)
    text(s, MX + Inches(0.7), y + Inches(0.1), Inches(6.0), Inches(0.8), st, size=16, color=INK)
code(s, MX + Inches(7.2), TOP, CW - Inches(7.2), Inches(4.6), [
    "# 常用參數：setup-host.cmd -參數",
    "",
    ">-Preset small|laptop|quality",
    "# 模型組合；預設依顯存自動選",
    "",
    ">-AllowFrom LocalSubnet,<VPN網段>",
    "# 防火牆開放來源（見第 05 段）",
    "",
    ">-SkipVllm",
    "# 只重做 Windows 端：防火牆、電源、opencode",
    "",
    ">-EnableSsh",
    "# 開 OpenSSH，讓人 ssh 進來用",
], size=13)

# =============================================================================
#  8. setup-host 做的事
# =============================================================================
s = content_slide("02 筆電主機安裝", "setup-host.cmd 自動完成的事",
                  "這八件事都是自動的，而且可以重複執行：再跑一次只會更新，不會蓋掉你改過的設定。"
                  "RTX 50 系列顯卡的相容設定也會自動處理，不需要另外裝 CUDA Toolkit。")
todo = ["檢查 NVIDIA 驅動、Windows 版本",
        "偵測 / 安裝 WSL2 Ubuntu",
        "WSL 改 mirrored 網路，和筆電共用 IP",
        "排程 vLLM-WSL-KeepAlive：登入後讓 WSL 常駐",
        "在 WSL 裝 vLLM、依顯存選模型、產生 API 金鑰",
        "防火牆開放 TCP 8000（預設只允許同網段）",
        "電源：插電時不睡眠、闔上螢幕不動作",
        "安裝 opencode，連到本機 vLLM"]
cw, rh = (CW - Inches(0.3)) / 2, Inches(0.95)
for i, t in enumerate(todo):
    col, row = i // 4, i % 4
    x, y = MX + (cw + Inches(0.3)) * col, TOP + (rh + Inches(0.15)) * row
    box(s, x, y, cw, rh)
    badge(s, x + Inches(0.25), y + Inches(0.2), Inches(0.55), str(i + 1), size=14)
    text(s, x + Inches(1.0), y, cw - Inches(1.2), rh, t, size=16, color=INK, anchor=MSO_ANCHOR.MIDDLE, space=0)

# =============================================================================
#  8b. 安裝完成：查看網址與金鑰
# =============================================================================
s = content_slide("02 筆電主機安裝 · 完成", "安裝完成：查看網址與 API 金鑰",
                  "安裝結束時畫面會印一次網址與金鑰，之後忘了隨時可以查。最簡單是雙擊 vllmctl.cmd 輸入 info；"
                  "在 PowerShell 也可以直接加參數。只想看金鑰的話，可以直接讀 WSL 裡的設定檔。"
                  "金鑰等同密碼，只私下給實驗室成員；外流就用 new-key 換一組。Discord bot 會自動讀，不用給它。")
code(s, MX, TOP, Inches(6.3), Inches(4.6), [
    "# 方法 1：雙擊 server\\wsl2\\vllmctl.cmd，輸入",
    ">info",
    "",
    "# 方法 2：PowerShell 直接執行（在 repo 資料夾）",
    ">server\\wsl2\\vllmctl.cmd info",
    "",
    "# 方法 3：只看金鑰",
    ">wsl -d Ubuntu -- grep SERVE_API_KEY ~/vllm-server/vllm.env",
    "",
    "# 換一組新金鑰（舊的立即失效）",
    ">server\\wsl2\\vllmctl.cmd new-key",
], size=12)   # 12pt：方法 3 的指令才能放在同一行，方便照抄
rx, rw = MX + Inches(6.6), CW - Inches(6.6)
text(s, rx, TOP, rw, Inches(0.35), "info 會印出", size=14, color=INK, bold=True)
box(s, rx, TOP + Inches(0.4), rw, Inches(1.55), fill=CODEBG, line=None, radius=0.04)
text(s, rx + Inches(0.25), TOP + Inches(0.6), rw - Inches(0.5), Inches(1.3),
     [[("模型名稱：qwen3-4b", {"color": CODETX})],
      [("vLLM 網址：http://192.168.1.23:8000/v1", {"color": ORANGE, "bold": True})],
      [("API 金鑰 ：sk-vllm-3f9c…", {"color": ORANGE, "bold": True})]], size=13, mono=True, space=4)
for i, (ic, t, d) in enumerate([("👥", "給同學：網址 + 金鑰", "同學雙擊 install-client.cmd 貼上"),
                                 ("💬", "Discord bot 不用給", "自動從 WSL 讀，換金鑰也會自動重讀"),
                                 ("🌐", "國網 API 金鑰是另一把", "來自國網的通知信，不在 vllmctl 裡")]):
    y = TOP + Inches(2.15) + Inches(0.83) * i
    box(s, rx, y, rw, Inches(0.72))
    badge(s, rx + Inches(0.15), y + Inches(0.11), Inches(0.5), ic, fill=TINT, size=13)
    text(s, rx + Inches(0.8), y + Inches(0.08), rw - Inches(0.95), Inches(0.6),
         [[(t, {"bold": True, "color": INK, "size": 14})], [(d, {"size": 12})]], space=1)

# =============================================================================
#  8c. 測試 vLLM：送一句話並計時
# =============================================================================
s = content_slide("02 主機安裝 · 測試（筆電與 Ubuntu 伺服器都適用）", "測試 vLLM：送一句話並計時",
                  "裝完、換模型、重開機之後，先在主機上自己測一次，再請同學連。指令都在 Linux 裡執行："
                  "筆電主機先輸入 wsl -d Ubuntu 進入 Ubuntu，Ubuntu 伺服器直接執行。"
                  "健康檢查印出 ok 代表服務在跑；計時測試會真的請模型回一句話，金鑰從 vllm.env 讀，不用貼。"
                  "正常 1 秒內回來；Couldn't connect 代表還在啟動或沒在跑；超過 10 秒或卡住，看下一頁的排查。"
                  "bench 是完整測速，會測單人、多人同時與長文件，約數分鐘。")
cw_ = Inches(6.55)
code(s, MX, TOP, cw_, Inches(4.6), [
    "# 筆電主機先進 Ubuntu（伺服器略過這行）",
    ">wsl -d Ubuntu",
    "",
    "# 1. 健康檢查：印出 ok 就是服務在跑",
    ">curl -sS -m 3 http://127.0.0.1:8000/health && echo ok",
    "",
    "# 2. 送一句話並計時（金鑰從 vllm.env 讀）",
    ">source ~/vllm-server/vllm.env; time curl -sS \\",
    ">  http://127.0.0.1:8000/v1/chat/completions \\",
    ">  -H \"Authorization: Bearer $SERVE_API_KEY\" \\",
    ">  -H \"Content-Type: application/json\" \\",
    ">  -d \"{\\\"model\\\":\\\"$SERVE_NAME\\\",\\\"max_tokens\\\":50,",
    ">  \\\"messages\\\":[{\\\"role\\\":\\\"user\\\",\\\"content\\\":\\\"用一句話介紹台北\\\"}]}\"",
    "",
    "# 3. 完整測速（筆電：vllmctl.cmd bench）",
    ">bash ~/local-llm/server/vllm/vllmctl.sh bench",
], size=11)
rx, rw = MX + cw_ + Inches(0.3), CW - cw_ - Inches(0.3)
text(s, rx, TOP, rw, Inches(0.35), "實際畫面：H200 回一句話只要 0.2 秒", size=14, color=INK, bold=True)
pic = screenshot(s, rx, TOP + Inches(0.42), rw, "vllm-test-curl.png")
highlight(s, pic, 0.0, 0.31, 1.0, 0.16, "1")
highlight(s, pic, 0.0, 0.745, 0.23, 0.07, "2")
ty = pic.top + pic.height + Inches(0.25)
table(s, rx, ty, rw, [
    ["看到", "代表"],
    ["① content 有中文回答", "模型正常回應"],
    ["② real 1 秒以內", "速度正常（第一次可能 1~2 秒）"],
    ["Couldn't connect to server", "還在啟動或沒在跑 → 看 logs"],
    ["超過 10 秒、卡住不動", "看下一頁：很慢或卡住"],
], [2.6, 3.0], size=12, rowh=0.4)

# =============================================================================
#  8d. 很慢或卡住：基本排查
# =============================================================================
s = content_slide("02 主機安裝 · 排查（實例：H200 每秒 0.2 token）", "很慢或卡住：基本排查指令",
                  "照順序看：服務狀態與紀錄 → GPU 有沒有在動 → CPU 誰在忙 → 用 py-spy 看 vLLM 卡在哪一行 → 驅動有沒有記錄 GPU 錯誤。"
                  "實例：H200 生成每秒只有 0.2 token，GPU 使用率 0%，連 systemd 都停不掉。py-spy 顯示停在 cudaMemcpyAsync，"
                  "也就是在等 GPU 傳輸；dmesg 沒有 Xid、PCIe Gen5 x16 正常，重開機後 Initializing model 從 107 秒變 6 秒，恢復正常。"
                  "GPU 傳輸測速：~/vllm-server/.venv/bin/python 跑 torch 把 256MB 搬到 GPU，"
                  "正常一般記憶體 5~20 GB/s、鎖定記憶體 30~55 GB/s。重開後還是一樣，或 dmesg 有 Xid，就交給機房檢查 GPU。")
cw_ = Inches(7.0)
code(s, MX, TOP, cw_, Inches(4.45), [
    "# 1. 服務狀態、最近紀錄（筆電：vllmctl.cmd status / logs）",
    ">bash ~/local-llm/server/vllm/vllmctl.sh status",
    ">journalctl -u vllm -n 40 --no-pager",
    ">journalctl -u vllm -f          # 即時看，Ctrl+C 離開",
    "",
    "# 2. GPU 有沒有在動（顯存、使用率）",
    ">nvidia-smi --query-gpu=memory.used,utilization.gpu --format=csv",
    "",
    "# 3. CPU 誰在忙（vLLM 都 0% = 在等東西，不是在算）",
    ">ps aux --sort=-%cpu | head -8",
    "",
    "# 4. vLLM 卡在哪一行（py-spy 第一次要先裝）",
    ">~/.local/bin/uv pip install --python ~/vllm-server/.venv/bin/python py-spy",
    ">sudo ~/vllm-server/.venv/bin/py-spy dump --pid $(pgrep -f EngineCore)",
    "",
    "# 5. 驅動有沒有記錄 GPU 錯誤（有 Xid 就是 GPU 出錯）",
    ">sudo dmesg -T | grep -iE \"xid|nvrm\" | tail -20",
], size=10)
rx, rw = MX + cw_ + Inches(0.3), CW - cw_ - Inches(0.3)
table(s, rx, TOP, rw, [
    ["看到", "怎麼辦"],
    ["Couldn't connect、ECONNREFUSED", "還在啟動：journalctl -f 等 startup complete"],
    ["紀錄有 Error / Traceback", "往上找第一個 Error 才是原因"],
    ["顯存不足（OutOfMemory）", "set SERVE_MAX_LEN 調小，或換小一點的組合"],
    ["很慢、GPU 0%、Running: 1 reqs", "引擎卡住 → 用 py-spy 看卡在哪"],
    ["py-spy 停在 cudaMemcpyAsync", "在等 GPU 傳輸 → sudo reboot"],
    ["dmesg 有 Xid、重開還是一樣", "GPU 或虛擬化平台問題 → 交給機房"],
], [2.45, 2.4], size=11, rowh=0.62)
box(s, MX, TOP + Inches(4.6), CW, Inches(0.5), fill=TINT, line=None)
text(s, MX + Inches(0.25), TOP + Inches(4.6), CW - Inches(0.5), Inches(0.5),
     [[("實例 H200：", {"bold": True, "color": INK}),
       ("每秒 0.2 token、GPU 0% → py-spy 停在 cudaMemcpyAsync → dmesg 沒 Xid → 重開機，"
        "Initializing model 107 秒 → 6 秒，恢復正常", {})]],
     size=12, anchor=MSO_ANCHOR.MIDDLE, space=0)

# =============================================================================
#  9. 同學電腦安裝 opencode
# =============================================================================
s = content_slide("03 同學使用 opencode", "同學的電腦：雙擊 install-client.cmd",
                  "同學的電腦不需要顯卡。安裝程式會先測試能不能連到筆電，連不到會直接說原因，"
                  "常見是筆電休眠、IP 打錯或不在同一個網路。預設安裝桌面版，不會用終端機的同學也能直接用；"
                  "習慣終端機的同學加 -Terminal 另外安裝終端機版。")
steps = ["向管理者拿「網址」和「API 金鑰」（筆電 vllmctl.cmd 輸入 info）",
         "複製專案的整個 opencode\\client 資料夾到自己電腦",
         "雙擊 install-client.cmd，貼上網址與金鑰",
         "開始選單開「OpenCode」，開啟專案資料夾"]
for i, st in enumerate(steps):
    y = TOP + Inches(1.0) * i
    badge(s, MX, y + Inches(0.05), Inches(0.5), str(i + 1), size=14)
    text(s, MX + Inches(0.7), y + Inches(0.1), Inches(6.0), Inches(0.8), st, size=16, color=INK)
code(s, MX + Inches(7.2), TOP, CW - Inches(7.2), Inches(2.6), [
    "# 網址範例（結尾要有 /v1）",
    ">http://192.168.1.23:8000/v1",
    "",
    "# 腳本自動完成",
    "1. 測試連線  2. 安裝桌面版",
    "3. 寫入 opencode.json  4. 裝共用規則",
], size=13)
box(s, MX + Inches(7.2), TOP + Inches(2.85), CW - Inches(7.2), Inches(1.75))
text(s, MX + Inches(7.5), TOP + Inches(3.05), CW - Inches(7.8), Inches(1.5),
     [[("加入其他伺服器：不用重裝", {"bold": True, "color": INK, "size": 15})],
      [("install-client.cmd -Url <網址> -Key <金鑰> -Name server2 -SkipInstall", {"mono": True, "size": 12})],
      [("終端機版：加 -Terminal　·　Ubuntu：install-client.sh", {"size": 12})]], space=6)

# =============================================================================
#  9b. servers.json：管理者預先填好伺服器
# =============================================================================
s = content_slide("03 同學使用 opencode · 管理者", "servers.json：先填好，同學雙擊就加入",
                  "管理者把 opencode\\servers.example.json 複製成 servers.json，填好每台伺服器的網址和金鑰，"
                  "再把整個 opencode 資料夾私下給同學。同學雙擊 install-client.cmd 就會依序加入每一台，不用自己輸入。"
                  "最常見的錯誤是 name 用了中文：name 只能英數字，中文名稱請放在 label。"
                  "國網這種每個人有自己金鑰的，key 留空，執行時會問。servers.json 有金鑰，已在 .gitignore 裡，不會上傳。")
code(s, MX, TOP, Inches(5.7), Inches(4.6), [
    '{ "servers": [',
    '  { "name": "labgpt",',
    '    "label": "實驗室GPT",',
    '    "url": "http://140.110.146.54:8000/v1",',
    '    "key": "sk-vllm-…",',
    '    "default": true },',
    '  { "name": "nchc",',
    '    "label": "NCHC GenAI",',
    '    "url": "https://inner-medusa.genai.nchc.org.tw/v1",',
    '    "key": "",',
    '    "context": 32768,',
    '    "models": ["gpt-oss-120b", "GLM-5.2"] }',
    ']}',
], size=12)
table(s, MX + Inches(6.0), TOP, CW - Inches(6.0), [
    ["欄位", "怎麼填"],
    ["name", "英數字代號，不能用中文、不能重複"],
    ["label", "opencode 裡顯示的名稱，可以用中文"],
    ["url", "伺服器網址，結尾要有 /v1"],
    ["key", "API 金鑰；留空 = 執行時再問（國網）"],
    ["default", "true = 設成預設模型（只設一台）"],
    ["models", "只加入這些模型；不填 = 全部"],
    ["context", "伺服器沒回報時的上下文（國網 32768）"],
    ["enabled", "false = 先不加入"],
], [1.5, 4.2], size=13, rowh=0.44, mono_cols=(0,))
text(s, MX + Inches(6.0), TOP + Inches(4.1), CW - Inches(6.0), Inches(0.5),
     [[("⚠️ servers.json 有金鑰：只私下傳給同學，不要上傳或貼到群組", {"color": RED, "bold": True})]], size=12)

# =============================================================================
#  9b-2. 實際畫面：install-client 依 servers.json 加入兩台
# =============================================================================
s = content_slide("03 同學使用 opencode · 實際畫面", "實際執行：servers.json 有兩台，依序加入",
                  "安裝前可以先用瀏覽器開伺服器網址（結尾 /v1）測試：看到 Unauthorized 代表連得到，只是瀏覽器沒帶金鑰；"
                  "轉圈圈或打不開才是連不到（IP 錯、沒開機、沒連 VPN）。"
                  "這是 servers.json 填了兩台伺服器時，雙擊 install-client.cmd 的實際畫面（右圖接在左圖後面）。"
                  "第一台實驗室 GPT：先測試連線並列出模型，接著安裝桌面版、寫入設定、裝共用規則，每一步都是 [OK]。"
                  "第二台國網的金鑰沒填對，出現 401 未經授權：這台會被跳過，其他照常安裝，修好金鑰後再雙擊一次就會補加入，"
                  "已經加過的不會重複。最後一步裝好文件工具 labdoc 和 /slides、/pdf2word 等指令，"
                  "看到「完成」後從開始選單開 OpenCode；OpenCode 原本開著的話要先關掉再開，才會讀到新設定。")
pw = (CW - Inches(0.3)) / 2
left = screenshot(s, MX, TOP, pw, "install-client-1.png")
right = screenshot(s, MX + pw + Inches(0.3), TOP, pw, "install-client-2.png")
highlight(s, left, 0.0, 0.15, 0.44, 0.095, "1")
highlight(s, right, 0.0, 0.445, 0.83, 0.125, "2")
highlight(s, right, 0.0, 0.595, 1.0, 0.175, "3")
cy = TOP + max(left.height, right.height) + Inches(0.25)
# 左下：安裝前先用瀏覽器測連線
bw = Inches(2.45)
browser = screenshot(s, MX, cy, bw, "browser-unauthorized.png")
highlight(s, browser, 0.0, 0.66, 0.43, 0.2, "0")
bx = MX + bw + Inches(0.2)
text(s, bx, cy, Inches(2.0), Inches(1.0),
     [[("安裝前：瀏覽器測試", {"bold": True, "color": INK, "size": 14})],
      [("開伺服器網址，看到 Unauthorized 就是連得到（只是沒帶金鑰）", {"size": 12})]], space=2, line=1.1)
x0 = bx + Inches(2.2)
cw = (MX + CW - x0 - Inches(0.4)) / 3
for i, (n, t, d) in enumerate([("1", "測試連線", "列出伺服器上的模型，[OK] 才往下裝"),
                               ("2", "401 未經授權", "金鑰沒填或填錯：先跳過，修好再雙擊"),
                               ("3", "文件工具與 /指令", "labdoc、/slides、/pdf2word 等一起裝好")]):
    x = x0 + (cw + Inches(0.2)) * i
    badge(s, x, cy + Inches(0.02), Inches(0.34), n, fill=ORANGE, size=12)
    text(s, x + Inches(0.48), cy, cw - Inches(0.5), Inches(0.9),
         [[(t, {"bold": True, "color": INK, "size": 14})], [(d, {"size": 12})]], space=2, line=1.1)

# =============================================================================
#  9c. 查看伺服器上有哪些模型
# =============================================================================
s = content_slide("03 同學使用 opencode · 管理者", "查看伺服器上有哪些模型",
                  "填 servers.json 的 models 欄位前，先查伺服器上有哪些模型。實驗室的 vLLM 一次只載入一個模型，"
                  "所以只會列出一個；國網會列出很多個，請挑支援工具呼叫的填進去。"
                  "PowerShell 的 Invoke-RestMethod 可以只列出模型名稱，方便直接複製。install-client 執行時第 1 步也會列出。")
code(s, MX, TOP, Inches(7.4), Inches(4.6), [
    "# 實驗室主機（在主機筆電上）",
    ">server\\wsl2\\vllmctl.cmd info      # 目前載入的模型",
    ">server\\wsl2\\vllmctl.cmd models    # 用過的模型（可 switch）",
    "",
    "# 任何電腦（PowerShell）：先設網址與金鑰",
    '>$u = "http://140.110.146.54:8000/v1"',
    '>$k = "你的金鑰"',
    '>(irm "$u/models" -Headers @{Authorization="Bearer $k"}).data.id',
    "",
    "# 國網：換成國網網址與自己的金鑰，再執行上面那行 irm",
    '>$u = "https://inner-medusa.genai.nchc.org.tw/v1"',
], size=12)
rx, rw = MX + Inches(7.7), CW - Inches(7.7)
for i, (ic, t, d) in enumerate([("1️⃣", "實驗室 vLLM：只會有 1 個", "一次只載入一個模型，換模型用 vllmctl"),
                                 ("📋", "國網：會列出很多個", "挑支援工具呼叫的，填進 servers.json 的 models"),
                                 ("🔎", "install-client 也會列出", "執行時第 1 步「測試連線」會顯示模型清單")]):
    y = TOP + Inches(1.55) * i
    box(s, rx, y, rw, Inches(1.35))
    badge(s, rx + Inches(0.2), y + Inches(0.25), Inches(0.6), ic, fill=TINT, size=16)
    text(s, rx + Inches(0.95), y + Inches(0.2), rw - Inches(1.1), Inches(1.0),
         [[(t, {"bold": True, "color": INK, "size": 15})], [(d, {"size": 13})]], space=3, line=1.1)

# =============================================================================
#  10. 第一次開啟 opencode
# =============================================================================
s = content_slide("03 同學使用 opencode", "第一次開啟 opencode",
                  "opencode 只會讀寫你開啟的那個專案資料夾底下的檔案，所以一定要先開啟專案資料夾。"
                  "第一次進專案先打 /init，它會產生 AGENTS.md 記錄專案結構，小模型有了它表現好很多。")
code(s, MX, TOP, Inches(5.6), Inches(4.4), [
    "# 桌面版",
    ">開始選單 → OpenCode → 開啟專案資料夾",
    "",
    "# 終端機版（-Terminal 才有）",
    "cd C:\\Users\\你\\Documents\\my-thesis",
    ">opencode",
    "",
    "# 第一次進這個專案",
    ">/init",
], size=15)
tips = [("下方輸入框", "用中文描述需求，Enter 送出"),
        ("模型選單 / 輸入 /models", "選 laptop/qwen3-4b = 筆電的 vLLM"),
        ("Tab 鍵", "切換 Build（會改檔）／ Plan（只分析不改）"),
        ("/init", "產生 AGENTS.md，之後每次對話都會參考")]
rx = MX + Inches(6.0)
for i, (t, d) in enumerate(tips):
    y = TOP + Inches(1.12) * i
    box(s, rx, y, CW - Inches(6.0), Inches(0.95))
    text(s, rx + Inches(0.3), y + Inches(0.12), CW - Inches(6.6), Inches(0.75),
         [[(t, {"bold": True, "color": INK, "size": 15})], [(d, {"size": 13})]], space=2)

# =============================================================================
#  11. 常用指令
# =============================================================================
s = content_slide("03 同學使用 opencode", "常用指令",
                  "打斜線會跳出完整的指令清單。最常用的是 /models、/compact 和 /undo。"
                  "@ 後面接檔名可以直接指定檔案，模型不用自己找，也比較省上下文。")
table(s, MX, TOP, CW, [
    ["輸入", "作用"],
    ["直接打字 + Enter", "用中文描述需求（回答會是繁體中文）"],
    ["@檔名", "指定檔案，例：@src/train.py"],
    ["Tab", "切換 Build（會改檔）／ Plan（只分析不改）"],
    ["/init", "分析專案、產生 AGENTS.md"],
    ["/models", "切換模型"],
    ["/undo · /redo", "復原 / 重做上一次修改"],
    ["/new · /compact", "開新對話 / 壓縮過長的對話（超過上下文時用）"],
    ["/exit 或 Ctrl+C", "離開"],
], [3, 7], size=15, rowh=0.5, mono_cols=(0,))

# =============================================================================
#  12. 提問技巧
# =============================================================================
s = content_slide("03 同學使用 opencode", "讓小模型也好用的提問技巧",
                  "4B 小模型可以做事，但要給明確、範圍小的任務。好的提問包含三件事：哪個檔案、哪個函式、要達到什麼結果。")
tips = [("1", "一次只做一件事", "「在 utils.py 加一個讀 CSV 的函式」比「重構整個專案」可靠得多。"),
        ("2", "直接指名檔案", "用 @src/train.py 指定檔案，寫出相對路徑，避免寫到 C:\\。"),
        ("3", "先 Plan 再 Build", "Tab 切到 Plan 讓它先列步驟，確認方向正確再切回 Build。"),
        ("4", "對話變長就重來", "回應變慢或出錯時用 /compact 或 /new；用中文問，回答就是中文。")]
cw, ch = (CW - Inches(0.3)) / 2, Inches(1.55)
for i, (n, t, d) in enumerate(tips):
    x, y = MX + (cw + Inches(0.3)) * (i % 2), TOP + (ch + Inches(0.25)) * (i // 2)
    box(s, x, y, cw, ch)
    badge(s, x + Inches(0.3), y + Inches(0.3), Inches(0.55), n, size=15)
    text(s, x + Inches(1.1), y + Inches(0.28), cw - Inches(1.4), Inches(1.1),
         [[(t, {"bold": True, "color": INK, "size": 17})], [(d, {"size": 14})]], space=4, line=1.1)
code(s, MX, TOP + Inches(3.65), CW, Inches(0.95), [
    "# 範例",
    "@data/loader.py 的 load_csv 讀到空白列會出錯，請修正並加上單元測試",
], size=14)

# =============================================================================
#  13. Discord bot 安裝
# =============================================================================
s = content_slide("04 Discord bot", "Discord bot：接同一台 vLLM",
                  "Discord bot 不另外架模型，直接用 opencode 的 vLLM。它會自動讀 WSL 裡的金鑰，"
                  "也會自動問 vLLM 目前是哪個模型，所以之後換模型、換金鑰，bot 都不用改設定。")
steps = ["先完成 02 的主機安裝（vLLM 要在跑）",
         "在 discord\\client 執行 deploy.ps1：建環境、自動填好 .env",
         "Developer Portal 建 bot：只開 MESSAGE CONTENT INTENT（下一頁）",
         "用 OAuth2 網址邀請 bot（bot + 四項權限），token 填進 .env",
         "測試：雙擊 start-bot.bat；正式：系統管理員執行 install-autostart.ps1"]
for i, st in enumerate(steps):
    y = TOP + Inches(0.92) * i
    badge(s, MX, y + Inches(0.05), Inches(0.5), str(i + 1), size=14)
    text(s, MX + Inches(0.7), y + Inches(0.1), Inches(6.2), Inches(0.8), st, size=16, color=INK)
code(s, MX + Inches(7.3), TOP, CW - Inches(7.3), Inches(2.3), [
    ">cd discord\\client",
    ">powershell -ExecutionPolicy Bypass",
    ">  -File .\\deploy.ps1",
    "",
    "# .env 只需要手動填這一行",
    "DISCORD_BOT_TOKEN=...",
], size=13)
for i, (t, d) in enumerate([("金鑰", "LLM_API_KEY 留空 → 自動從 WSL 讀，換金鑰也會自動重讀"),
                            ("模型", "自動用 vLLM 目前的模型，Discord 輸入 !model 可確認")]):
    y = TOP + Inches(2.5) + Inches(1.1) * i
    box(s, MX + Inches(7.3), y, CW - Inches(7.3), Inches(0.95))
    text(s, MX + Inches(7.6), y + Inches(0.12), CW - Inches(7.9), Inches(0.75),
         [[(t, {"bold": True, "color": INK, "size": 15})], [(d, {"size": 13})]], space=2)

# =============================================================================
#  13b. Developer Portal 設定
# =============================================================================
s = content_slide("04 Discord bot · 管理者", "Developer Portal：只開需要的權限",
                  "程式只用到 Message Content 這一個 privileged intent，其他兩個關掉，少收一類事件也少一份個資。"
                  "公開機器人一定要關：開著的話，邀請連結外流時別人也能把 bot 拉進他們的伺服器，用的是實驗室的顯卡。"
                  "關掉之後只有應用程式擁有者能邀請。重新導向 URI 只邀 bot 用不到，可以不填。")
table(s, MX, TOP, Inches(7.4), [
    ["Bot 頁的開關", "建議", "原因"],
    ["公開機器人 Public Bot", "關", "開著時，拿到邀請連結的人都能把 bot 加進自己的伺服器"],
    ["需要 OAuth2 代碼授權", "關", "只邀請 bot 用不到"],
    ["Presence Intent", "關", "程式沒用到"],
    ["Server Members Intent", "關", "程式沒用到；之後要依身分組限制才開"],
    ["Message Content Intent", "開", "必要：! 指令要讀訊息內容"],
], [2.5, 0.7, 4.2], size=13, rowh=0.62)
rx, rw = MX + Inches(7.7), CW - Inches(7.7)
box(s, rx, TOP, rw, Inches(2.6))
text(s, rx + Inches(0.3), TOP + Inches(0.25), rw - Inches(0.6), Inches(2.2),
     [[("OAuth2 → URL 產生器", {"bold": True, "color": INK, "size": 16})],
      [("範圍：", {"size": 13}), ("bot", {"size": 13, "mono": True, "color": INK})],
      [("權限：", {"size": 13}),
       ("View Channels · Send Messages · Read Message History · Attach Files", {"size": 13, "mono": True, "color": INK})],
      [("Attach Files 用來回傳 pptx / docx / xlsx；不要勾 Administrator。重新導向 URI 不用填。", {"size": 13})]],
     space=8, line=1.1)
box(s, rx, TOP + Inches(2.8), rw, Inches(1.25))
text(s, rx + Inches(0.3), TOP + Inches(2.97), rw - Inches(0.6), Inches(1.0),
     [[("⚠️ Token 等於 bot 本身", {"bold": True, "color": RED, "size": 15})],
      [("只放 .env；不要截圖、貼 Notion 或 Discord。外洩就到 Bot 頁 Reset Token", {"size": 13})]],
     space=4, line=1.1)


def captions(s, y, items, x=MX, w=CW):
    """截圖下方的編號說明列：[(編號, 標題, 說明), ...]"""
    cw = (w - Inches(0.25) * (len(items) - 1)) / len(items)
    for i, (n, t, d) in enumerate(items):
        cx = x + (cw + Inches(0.25)) * i
        badge(s, cx, y + Inches(0.02), Inches(0.34), n, fill=ORANGE, size=12)
        text(s, cx + Inches(0.48), y, cw - Inches(0.5), Inches(1.0),
             [[(t, {"bold": True, "color": INK, "size": 14})], [(d, {"size": 12})]], space=2, line=1.1)


# =============================================================================
#  13a. 實際畫面：建立 Application、拿 token
# =============================================================================
s = content_slide("04 Discord bot · 管理者 · 實際畫面", "建立 bot：新建應用程式，到「機器人」頁拿 token",
                  "到 discord.com/developers/applications 按「新建應用程式」，名稱取好認的，例如 LabGPT-H200。"
                  "每台要跑 bot 的主機各建一個：同一個 token 同時在兩台跑，每則訊息會被回兩次。"
                  "建好後左邊選「機器人」，按「重設權杖」再按「複製」。token 只會顯示這一次，"
                  "馬上貼到 deploy.sh --token 或 .env；之後忘了就再重設一次，舊的會立刻失效。"
                  "這個 token 是 Discord 的，不是 vLLM 的 sk-vllm- 金鑰，vLLM 金鑰 bot 會自動讀。")
pw = (CW - Inches(0.3)) / 2
left = screenshot(s, MX, TOP, pw, "discord-1-apps.png")
right = screenshot(s, MX + pw + Inches(0.3), TOP, pw, "discord-2-token.png")
highlight(s, left, 0.855, 0.2, 0.125, 0.075, "1")
highlight(s, right, 0.008, 0.395, 0.125, 0.048, "2")
highlight(s, right, 0.16, 0.785, 0.45, 0.2, "3")
captions(s, TOP + max(left.height, right.height) + Inches(0.3), [
    ("1", "新建應用程式", "名稱例如 LabGPT-H200；每台主機各建一個"),
    ("2", "左邊選「機器人」", "token 和下一頁的開關都在這一頁"),
    ("3", "重設權杖 → 複製", "只顯示一次；不是 vLLM 的 sk-vllm- 金鑰"),
])

# =============================================================================
#  13b. 實際畫面：機器人頁的開關
# =============================================================================
s = content_slide("04 Discord bot · 管理者 · 實際畫面", "機器人頁的開關：圖中是改之前的狀態",
                  "這是剛建好 bot 時的預設狀態，三個地方要改。公開機器人要關，不然邀請連結外流時別人也能把 bot 拉進他們的伺服器。"
                  "Presence 和 Server Members 程式沒用到，關掉。Message Content 一定要開：圖中沒開，"
                  "bot 啟動時就會出現 PrivilegedIntentsRequired 而一直重啟，Discord 上看到 bot 是離線。"
                  "改完要按畫面下方跳出來的「儲存變更」。改開關不用重設 token；存好後 bot 會在 10 秒內自動重試連上。")
pw = Inches(7.7)
pic = screenshot(s, MX, TOP, pw, "discord-3-intents.png")
highlight(s, pic, 0.16, 0.265, 0.83, 0.072, "1")
highlight(s, pic, 0.16, 0.625, 0.83, 0.22, "2")
highlight(s, pic, 0.16, 0.882, 0.83, 0.09, "3")
rx, rw = MX + pw + Inches(0.3), CW - pw - Inches(0.3)
for i, (n, t, d, col) in enumerate([
        ("1", "公開機器人 → 關", "圖中開著：邀請連結外流，別人也能把 bot 拉走，用實驗室的顯卡", INK),
        ("2", "Presence、Server Members → 關", "程式沒用到，少收一類事件也少一份個資", INK),
        ("3", "Message Content → 開", "圖中沒開：bot 起不來（PrivilegedIntentsRequired）", RED)]):
    y = TOP + Inches(1.05) * i
    box(s, rx, y, rw, Inches(0.92))
    badge(s, rx + Inches(0.15), y + Inches(0.15), Inches(0.34), n, fill=ORANGE, size=12)
    text(s, rx + Inches(0.62), y + Inches(0.1), rw - Inches(0.75), Inches(0.8),
         [[(t, {"bold": True, "color": col, "size": 14})], [(d, {"size": 11})]], space=2, line=1.05)
box(s, rx, TOP + Inches(3.2), rw, Inches(0.85), fill=TINT, line=None)
text(s, rx + Inches(0.2), TOP + Inches(3.2), rw - Inches(0.4), Inches(0.85),
     [[("改完按「儲存變更」", {"bold": True, "color": INK, "size": 14})],
      [("不用重設 token，bot 10 秒內會自動重試", {"size": 12})]], space=2, anchor=MSO_ANCHOR.MIDDLE)

# =============================================================================
#  13c. 實際畫面：OAuth2 產生邀請網址、授權加入伺服器
# =============================================================================
s = content_slide("04 Discord bot · 管理者 · 實際畫面", "OAuth2：產生邀請網址，把 bot 加進伺服器",
                  "左邊選 OAuth2，往下到「OAuth2 URL 產生器」：範圍只勾 bot，下方的機器人權限勾檢視頻道、傳送訊息、"
                  "附加檔案、讀取訊息歷史記錄這四個（圖中已捲到下面，傳送訊息在上方）。重新導向不用填。"
                  "整合類型選公會安裝，複製最下面產生的網址，用瀏覽器打開，選要加入的伺服器、按繼續再按授權。"
                  "要有那個伺服器的「管理伺服器」權限才看得到它。機器人頁也有一個權限計算器，那個不會產生邀請網址，用這一頁就好。")
lw = Inches(6.2)
left = screenshot(s, MX, TOP, lw, "discord-4-oauth2-url.png")
right = screenshot(s, MX + lw + Inches(0.3), TOP, CW - lw - Inches(0.3), "discord-5-authorize.png")
highlight(s, left, 0.185, 0.41, 0.075, 0.05, "1")
highlight(s, left, 0.448, 0.255, 0.12, 0.09, "1")
highlight(s, left, 0.175, 0.905, 0.81, 0.075, "2")
highlight(s, right, 0.49, 0.75, 0.24, 0.06, "3")
highlight(s, right, 0.61, 0.885, 0.165, 0.065, "4")
captions(s, TOP + max(left.height, right.height) + Inches(0.3), [
    ("1", "範圍 bot ＋ 四個權限", "檢視頻道、傳送訊息、附加檔案、讀取訊息歷史記錄"),
    ("2", "複製產生的 URL", "用瀏覽器打開；重新導向不用填"),
    ("3", "選要加入的伺服器", "要有「管理伺服器」權限才看得到"),
    ("4", "繼續 → 授權", "bot 出現在成員列表（還沒啟動前是離線）"),
])

# =============================================================================
#  13d. 實際畫面：Ubuntu 伺服器部署 bot
# =============================================================================
s = content_slide("04 Discord bot · 管理者 · 實際畫面", "在伺服器部署 bot：deploy.sh 一行完成（Ubuntu / H200）",
                  "bot 跟 vLLM 跑在同一台，金鑰自動讀 ~/vllm-server/vllm.env，不用貼。deploy.sh 會確認 vLLM、建 Python 環境、"
                  "產生 .env，最後註冊成 systemd 服務 discord-bot：開機自動啟動、掛掉 10 秒後自動重啟。"
                  "第 4 步要輸入 Ubuntu 密碼（sudo）。看到「[已上線]」就完成，在 Discord 輸入 !help 試試看。"
                  "Windows 筆電改用 discord\\client\\deploy.ps1（vLLM 在 WSL），步驟見 discord-bot.md。")
iw = Inches(3.25)
pic = screenshot(s, MX + CW - iw, TOP - Inches(0.2), iw, "discord-6-deploy.png")
highlight(s, pic, 0.0, 0.355, 1.0, 0.03, "1")
highlight(s, pic, 0.0, 0.868, 1.0, 0.045, "2")
highlight(s, pic, 0.0, 0.93, 1.0, 0.065, "3")
lw = CW - iw - Inches(0.3)
code(s, MX, TOP, lw, Inches(3.05), [
    "# 1. 拉最新程式，部署（token 是 Discord 的，不是 sk-vllm-）",
    ">cd ~/local-llm && git pull",
    ">bash discord/client/deploy.sh --token <Discord_bot_token>",
    "",
    "# 2. 看紀錄：出現 [已上線] 就完成（Ctrl+C 離開）",
    ">journalctl -u discord-bot -f",
    "",
    "# 日常：狀態 / 改完 .env 重啟 / 移除",
    ">sudo systemctl status discord-bot",
    ">sudo systemctl restart discord-bot",
    ">bash discord/client/deploy.sh --uninstall",
], size=11)
table(s, MX, TOP + Inches(3.2), lw, [
    ["紀錄出現", "處理"],
    ["① vLLM 正常 ② 已建立 .env ③ 註冊服務", "deploy.sh 三個重點步驟（右圖）"],
    ["PrivilegedIntentsRequired", "開 Message Content Intent 並儲存，不用重設 token"],
    ["LoginFailure", "token 貼錯：重新複製填進 .env，再 restart"],
    ["已上線但回很慢", "vLLM 卡住了：用「很慢或卡住」那頁的指令排查"],
], [3.0, 4.6], size=11, rowh=0.37)

# =============================================================================
#  14. Discord 指令
# =============================================================================
s = content_slide("04 Discord bot", "Discord 指令",
                  "產檔指令可以附上 txt、md、docx、pdf，或用 Discord 的「回覆」功能回覆某則訊息，模型會以那些內容為準。"
                  "例如回覆一則會議記錄，輸入 !ppt 整理成組會報告。產出的都是草稿，文獻一律標為待補，模型被要求不可以捏造。")
table(s, MX, TOP, Inches(7.4), [
    ["指令", "作用"],
    ["@bot / 私訊 / 直接發言", "一般問答"],
    ["!web <問題>", "上網搜尋後回答"],
    ["!ppt <主題>", "簡報 .pptx（含講者備註）"],
    ["!word 計畫 <題目>", "國科會研究計畫草稿 .docx"],
    ["!word 報告 <主題>", "課堂報告 .docx"],
    ["!excel <需求>", "表格 .xlsx（可含公式）"],
    ["!model · !reset · !help", "目前模型 · 清記憶 · 說明"],
], [3.2, 4.2], size=14, rowh=0.5, mono_cols=(0,))
rx, rw = MX + Inches(7.8), CW - Inches(7.8)
box(s, rx, TOP, rw, Inches(2.3))
text(s, rx + Inches(0.3), TOP + Inches(0.25), rw - Inches(0.6), Inches(1.9),
     [[("📎 附參考資料", {"bold": True, "color": INK, "size": 16})],
      [("附上 .txt / .md / .docx / .pdf，或「回覆」某則訊息再下指令，內容會以參考資料為準", {"size": 14})]],
     space=6, line=1.1)
box(s, rx, TOP + Inches(2.5), rw, Inches(1.5))
text(s, rx + Inches(0.3), TOP + Inches(2.7), rw - Inches(0.6), Inches(1.2),
     [[("⚠️ 產出是草稿", {"bold": True, "color": RED, "size": 16})],
      [("文獻標為 [待補]，數據與論述請自行查證", {"size": 14})]], space=6)

# =============================================================================
#  14b. 頻道與身分組
# =============================================================================
s = content_slide("04 Discord bot · 上線前", "頻道與身分組：先決定 bot 在哪裡說話",
                  "程式目前不檢查身分組，所以誰能用是靠 Discord 的頻道權限：bot 的身分組只開放 llm 頻道。"
                  "ALLOWED_CHANNEL_ID 預設是空的，這時 bot 會回覆它看得到的每個頻道的每則訊息，上線前一定要填 #llm-問答 的頻道 ID"
                  "（Discord 設定開啟開發者模式，對頻道按右鍵「複製頻道 ID」）。"
                  "私訊不受頻道權限管，只要和 bot 在同一個伺服器就能私訊它。")
table(s, MX, TOP, Inches(7.4), [
    ["頻道", "用途", "bot 權限"],
    ["#llm-問答", "一般問答；ID 填進 ALLOWED_CHANNEL_ID", "可讀、可發言"],
    ["#llm-文件", "!ppt / !word / !excel", "可讀、可發言、可附檔"],
    ["#llm-公告", "維護、換模型、離線通知", "唯讀"],
    ["#llm-回報", "錯誤與建議", "看不到"],
], [1.5, 3.6, 2.3], size=13, rowh=0.62, mono_cols=(0,))
rx, rw = MX + Inches(7.7), CW - Inches(7.7)
for i, (t, d) in enumerate([("@實驗室成員", "才看得到 llm 頻道"),
                            ("@LLM 管理者", "管 bot 設定、發公告"),
                            ("@llm-bot", "只給 llm 頻道的權限")]):
    y = TOP + Inches(1.05) * i
    box(s, rx, y, rw, Inches(0.9))
    text(s, rx + Inches(0.3), y + Inches(0.12), rw - Inches(0.6), Inches(0.7),
         [[(t, {"bold": True, "color": INK, "size": 15})], [(d, {"size": 13})]], space=2)
box(s, MX, TOP + Inches(3.4), CW, Inches(1.1), fill=TINT, line=None)
text(s, MX + Inches(0.3), TOP + Inches(3.55), CW - Inches(0.6), Inches(0.85),
     [[("⚠️ .env 的 ALLOWED_CHANNEL_ID 預設空白", {"bold": True, "color": RED, "size": 15})],
      [("空白時 bot 會回覆它看得到的每個頻道的每則訊息，上線前一定要填", {"size": 13})]], space=4)

# =============================================================================
#  14c. 安全與隱私
# =============================================================================
s = content_slide("04 Discord bot · 上線前", "安全與隱私：模型在本地，對話不在本地",
                  "opencode 的對話只在實驗室，但 Discord bot 的提問、附件和產出的檔案都會經過 Discord 的伺服器，"
                  "要跟同學講清楚。bot 的對話記憶是以頻道為單位存在記憶體裡，同頻道所有人共用，bot 重啟就清空。")
tips = [("1", "訊息存在 Discord 雲端", "提問、附件、產出的檔案都經過 Discord；未公開數據、個資、帳密不要貼"),
        ("2", "記憶是整個頻道共用", "同頻道的人共用最近 10 則對話；換話題先 !reset，私事用私訊"),
        ("3", "Token 只放 .env", ".env 已在 .gitignore；截圖、群組都別貼，外洩就 Reset Token"),
        ("4", "產出是草稿", "文獻標 [待補]、數據要查證，交出去前要人工檢查")]
cw, ch = (CW - Inches(0.3)) / 2, Inches(1.7)
for i, (n, t, d) in enumerate(tips):
    x, y = MX + (cw + Inches(0.3)) * (i % 2), TOP + (ch + Inches(0.3)) * (i // 2)
    box(s, x, y, cw, ch)
    badge(s, x + Inches(0.3), y + Inches(0.3), Inches(0.55), n, size=15)
    text(s, x + Inches(1.1), y + Inches(0.28), cw - Inches(1.4), Inches(1.5),
         [[(t, {"bold": True, "color": INK, "size": 17})], [(d, {"size": 14})]], space=4, line=1.1)

# =============================================================================
#  14d. 資源與維運
# =============================================================================
s = content_slide("04 Discord bot · 管理者", "資源與維運：一張顯卡，出事要查得到",
                  "bot 和 opencode 共用同一張顯卡，4B 模型五人同時的總吞吐約 121 tok/s。"
                  "程式目前沒有限流和排隊，一個人連續丟好幾個 !ppt 就會拖慢其他人。"
                  "右欄是建議，還沒寫進程式；log 只記中繼資料，不記訊息全文。")
table(s, MX, TOP, CW, [
    ["項目", "現況", "建議"],
    ["共用顯卡", "bot 與 opencode 共用，沒有限流", "每人冷卻、一次一個產檔；排隊時先回「處理中」"],
    ["筆電休眠", "睡眠或重開機時 bot 離線", "電源設永不睡眠；離線時在 #llm-公告 通知"],
    ["使用紀錄", "只在終端機印上線訊息", "寫 log：時間、使用者、指令、耗時、錯誤"],
    ["健康檢查", "同學自己發現沒回應", "加 !status：目前模型、排隊數"],
    ["Token / 套件", "沒有固定規則", "每學期 Reset Token；定期更新 discord.py"],
    ["交接", "設定只在個人帳號", "Developer Portal 建 Team，加第二位管理者"],
], [1.6, 3.6, 5.0], size=14, rowh=0.6)

# =============================================================================
#  14e. 上線前檢查與待決定
# =============================================================================
s = content_slide("04 Discord bot · 上線前", "上線前檢查，與要一起決定的事",
                  "左邊照順序做完就能上線。右邊四件事牽涉程式修改或使用規範，建議組會討論後再決定。"
                  "斜線指令是 Discord 目前推薦的做法，改用後 bot 就不需要讀所有訊息的內容。")
table(s, MX, TOP, Inches(6.9), [
    ["狀態", "項目"],
    ["完成", "建立 llm-bot，開啟 Message Content Intent"],
    ["待辦", "關閉公開機器人、Presence、Server Members"],
    ["待辦", "用 OAuth2 網址邀請 bot（bot + 四項權限）"],
    ["待辦", "建頻道與身分組，填 ALLOWED_CHANNEL_ID"],
    ["待辦", "deploy.ps1、填 Token、install-autostart.ps1"],
    ["待辦", "每個指令實測一次（含附件、回覆訊息）"],
    ["待辦", "重開機後確認 bot 自動上線"],
    ["待辦", "#llm-公告 貼使用規範與指令表"],
], [0.9, 6.0], size=13, rowh=0.5)
rx, rw = MX + Inches(7.2), CW - Inches(7.2)
for i, (t, d) in enumerate([("誰能用？", "全伺服器，或只限 @實驗室成員（要改程式檢查身分組）"),
                            ("開放私訊嗎？", "私訊不受頻道權限管，同伺服器的人都能用"),
                            ("記憶怎麼分？", "維持每頻道共用，或改成每人、每個討論串"),
                            ("改用斜線指令？", "/ppt 有提示好上手，長期可不用 Message Content")]):
    y = TOP + Inches(1.15) * i
    box(s, rx, y, rw, Inches(1.0))
    text(s, rx + Inches(0.3), y + Inches(0.12), rw - Inches(0.6), Inches(0.8),
         [[(t, {"bold": True, "color": INK, "size": 15})], [(d, {"size": 13})]], space=2, line=1.1)

# =============================================================================
#  15. 不在實驗室時
# =============================================================================
s = content_slide("05 不在實驗室時", "三種使用者，三種連線方式",
                  "實驗室同學在家時，連上學校 VPN 就等於在校內網路，其他設定完全不用改，也不用裝 Tailscale。"
                  "只有沒有學校帳號、不能用 VPN 的校外合作者，才需要用 Tailscale Funnel，而且用完就關。")
cols = [("🏫", "在實驗室", "實驗室同學", "http://筆電IP:8000/v1", ["區網直連", "什麼都不用多做"], BLUE),
        ("🏠", "在家", "實驗室同學", "先連學校 VPN", ["之後跟在實驗室一樣", "第一次要把 VPN 網段", "加進防火牆"], BLUE),
        ("🌐", "校外合作者", "沒有學校帳號", "https://筆電.xxx.ts.net/v1", ["Tailscale Funnel", "對方不用裝 Tailscale", "用完就 off"], ORANGE)]
cw = (CW - Inches(0.5)) / 3
for i, (ic, t, who, url, pts, col) in enumerate(cols):
    x = MX + (cw + Inches(0.25)) * i
    box(s, x, TOP, cw, Inches(3.75))
    badge(s, x + Inches(0.35), TOP + Inches(0.35), Inches(0.75), ic, fill=TINT, size=22)
    text(s, x + Inches(0.35), TOP + Inches(1.3), cw - Inches(0.7), Inches(0.9),
         [[(t, {"bold": True, "color": INK, "size": 22})], [(who, {"size": 13, "color": MUTED})]], space=2)
    text(s, x + Inches(0.35), TOP + Inches(2.2), cw - Inches(0.7), Inches(0.5), url, size=13, color=col, bold=True, mono=True)
    text(s, x + Inches(0.35), TOP + Inches(2.8), cw - Inches(0.7), Inches(1.4), pts, size=14, color=BODY, space=4)

# =============================================================================
#  16. 在家：學校 VPN
# =============================================================================
s = content_slide("05 不在實驗室時 · 在家", "在家：先連學校 VPN",
                  "第一次需要管理者和一位同學一起做，大約十分鐘。重點是找出 VPN 使用者拿到的網段，加進筆電的防火牆。"
                  "LocalSubnet 一定要保留，不然實驗室裡的同學會被擋。不要用 Any 偷懶，學校常配公開 IP，Any 等於整個網際網路都連得到。")
steps = ["同學在家連上 VPN，ipconfig 看 VPN 網卡的 IP",
         "瀏覽器開 http://筆電IP:8000/v1/models",
         "看到 Unauthorized → 已經通了，跳到第 5 步",
         "一直轉圈 → 問計中 VPN 網段，加進防火牆（右邊指令）",
         "雙擊 install-client.cmd，貼上網址與金鑰"]
for i, st in enumerate(steps):
    y = TOP + Inches(0.92) * i
    badge(s, MX, y + Inches(0.05), Inches(0.5), str(i + 1), size=14)
    text(s, MX + Inches(0.7), y + Inches(0.1), Inches(6.0), Inches(0.8), st, size=16, color=INK)
code(s, MX + Inches(7.0), TOP, CW - Inches(7.0), Inches(2.5), [
    "# 筆電，會要求系統管理員權限",
    ">server\\wsl2\\setup-host.cmd -SkipVllm",
    ">  -AllowFrom LocalSubnet,10.8.0.0/16",
    "",
    "# 10.8.0.0/16 換成學校 VPN 的網段",
    "# 多個網段用逗號接下去",
], size=13)
box(s, MX + Inches(7.0), TOP + Inches(2.7), CW - Inches(7.0), Inches(1.9))
text(s, MX + Inches(7.3), TOP + Inches(2.9), CW - Inches(7.6), Inches(1.6),
     [[("⚠️ 不要用 -AllowFrom Any", {"bold": True, "color": RED, "size": 15})],
      [("學校常配 140.x 公開 IP，Any = 整個網際網路都連得到 8000 port。", {"size": 13})],
      [("還是不通：學校 VPN 可能沒開放連實驗室網段，請計中協助。", {"size": 13})]], space=5, line=1.1)

# =============================================================================
#  17. 校外合作者：Funnel
# =============================================================================
s = content_slide("05 不在實驗室時 · 校外合作者", "校外合作者：Tailscale Funnel",
                  "Funnel 會把 vLLM 公開成一個 HTTPS 網址，對方不用裝任何東西。但這代表全世界都連得到這個網址，"
                  "唯一的保護是 API 金鑰，所以給之前先換金鑰，用完就關。腳本只公開 /v1，因為 vLLM 的金鑰只保護 /v1，"
                  "/metrics 這類路徑不需要金鑰。")
code(s, MX, TOP, Inches(5.6), Inches(4.4), [
    "# 筆電：雙擊 server\\wsl2\\funnel.cmd",
    ">on",
    "# 第一次：自動裝 Tailscale → 登入",
    "# → 瀏覽器按「同意啟用 Funnel」",
    "",
    ">test",
    "# 沒金鑰 401、有金鑰 200、/metrics 404",
    "",
    ">off",
    "# 不用時就關",
], size=14)
rx, rw = MX + Inches(6.0), CW - Inches(6.0)
for i, (ic, t, d) in enumerate([("🔑", "給之前先換金鑰", "vllmctl.cmd new-key；合作結束再換一次"),
                                 ("🛡️", "只公開 /v1", "HTTPS 由 Tailscale 自動處理"),
                                 ("⏱️", "第一次最多等 10 分鐘", "公開 DNS 生效；手機關 Wi-Fi 測，看到 401 即成功"),
                                 ("📉", "Beta 功能、有頻寬限制", "適合少數人臨時用")]):
    y = TOP + Inches(1.12) * i
    box(s, rx, y, rw, Inches(0.95))
    badge(s, rx + Inches(0.2), y + Inches(0.17), Inches(0.6), ic, fill=TINT, size=16)
    text(s, rx + Inches(1.0), y + Inches(0.12), rw - Inches(1.2), Inches(0.75),
         [[(t, {"bold": True, "color": INK, "size": 15})], [(d, {"size": 12})]], space=2)

# =============================================================================
#  18. vllmctl
# =============================================================================
s = content_slide("06 管理與排錯", "日常管理：雙擊 server\\wsl2\\vllmctl.cmd",
                  "vllmctl 只在筆電上用，雙擊後輸入指令。最常用的是 info 和 logs。"
                  "stop 可以暫時釋放顯卡，例如筆電要跑其他 GPU 程式時。")
table(s, MX, TOP, CW, [
    ["指令", "作用"],
    ["info", "給同學的網址、金鑰、安裝方式"],
    ["show", "目前設定（金鑰遮蔽）、GPU 用量、KV cache 容量"],
    ["status · logs", "服務狀態 / 即時紀錄（Ctrl+C 離開）"],
    ["start · stop · restart", "啟動 / 停止（釋放 GPU）/ 重啟"],
    ["presets · preset 名稱", "列出 / 換成內建模型組合"],
    ["model HF模型ID 名稱 上下文 parser", "換成任意 Hugging Face 模型"],
    ["models · switch 名稱", "列出用過的模型 / 切換回去（設定會保留）"],
    ["set 參數 值 · edit", "修改 vllm.env 參數並重啟"],
    ["new-key · bench", "產生新金鑰 / 跑標準測試"],
], [4, 6], size=14, rowh=0.46, mono_cols=(0,))

# =============================================================================
#  19. API 金鑰
# =============================================================================
s = content_slide("06 管理與排錯 · 金鑰", "API 金鑰：只有一組，大家共用",
                  "金鑰存在筆電 WSL 的 vllm.env。換金鑰後，Discord bot 會自動重讀，不用動；"
                  "但每台 opencode 都要用新金鑰重跑 install-client，筆電本機的 opencode 可以自動讀金鑰，不用貼。"
                  "Hugging Face token 是另一回事，只有下載需要授權的模型才要。")
for i, (t, cmd, d) in enumerate([("查看", "vllmctl.cmd → info", "要給同學時"),
                                   ("換新", "vllmctl.cmd → new-key", "有人離開、金鑰外流時；舊的立即失效")]):
    x = MX + (Inches(6.0) + Inches(0.3)) * i
    box(s, x, TOP, Inches(6.0), Inches(1.3))
    text(s, x + Inches(0.3), TOP + Inches(0.2), Inches(5.4), Inches(1.0),
         [[(t + "  ", {"bold": True, "color": INK, "size": 18}), (cmd, {"mono": True, "color": BLUE, "size": 15})],
          [(d, {"size": 14})]], space=6)
text(s, MX, TOP + Inches(1.6), CW, Inches(0.4), "換金鑰之後", size=16, color=INK, bold=True)
table(s, MX, TOP + Inches(2.05), CW, [
    ["誰", "要做什麼"],
    ["Discord bot", "✅ 不用動（遇到金鑰錯誤會自動從 WSL 重讀）"],
    ["筆電的 opencode", "install-client.cmd -Url http://127.0.0.1:8000/v1 -SkipInstall（自動讀金鑰）"],
    ["同學的 opencode", "收到新金鑰後，再雙擊一次 install-client.cmd"],
], [2.5, 7.5], size=14, rowh=0.5)
text(s, MX, TOP + Inches(4.25), CW, Inches(0.35),
     "Hugging Face token（hf_ 開頭）只有下載 Llama、Gemma 等需授權模型才需要：vllmctl.cmd → set HF_TOKEN hf_xxx",
     size=12, color=MUTED)

# =============================================================================
#  20. 換模型
# =============================================================================
s = content_slide("06 管理與排錯 · 換模型", "vLLM 換模型：五個步驟",
                  "vLLM 一次只載入一個模型，換模型就是下載新的再換掉舊的。vllmctl 會先下載完才切換，"
                  "下載期間大家照常使用，切換時中斷約一分鐘。換完一定要跑 bench；不滿意就用 switch 換回來，"
                  "設定會保留、也不用重新下載。最後請大家重跑 install-client，預設模型會自動更新。Discord bot 不用動。")
flow = [("1", "換模型", "preset quality\n或 model …"),
        ("2", "跑測試", "bench\n（約 1 分鐘）"),
        ("3", "不滿意就換回", "switch\nqwen3-4b"),
        ("4", "更新 opencode", "每台重跑\ninstall-client"),
        ("5", "Discord bot", "不用動\n!model 確認")]
fw = (CW - Inches(0.25) * 4) / 5
for i, (n, t, d) in enumerate(flow):
    x = MX + (fw + Inches(0.25)) * i
    box(s, x, TOP, fw, Inches(1.75))
    badge(s, x + Inches(0.2), TOP + Inches(0.22), Inches(0.48), n, size=14, fill=ORANGE if i == 4 else BLUE)
    text(s, x + Inches(0.78), TOP + Inches(0.25), fw - Inches(0.85), Inches(0.45), t, size=15, color=INK, bold=True)
    text(s, x + Inches(0.2), TOP + Inches(0.85), fw - Inches(0.35), Inches(0.85), d.split("\n"), size=13, space=1, mono=True)
code(s, MX, TOP + Inches(2.05), Inches(5.6), Inches(2.55), [
    "# 都在 server\\wsl2\\vllmctl.cmd 裡輸入",
    ">preset quality   # 內建組合：Qwen3-8B",
    ">model Qwen/Qwen3-8B-AWQ qwen3-8b 32768 hermes",
    ">models           # 列出用過的模型",
    ">switch qwen3-4b  # 換回去（設定保留）",
    ">bench            # 跑標準測試",
], size=13)
table(s, MX + Inches(5.9), TOP + Inches(2.05), CW - Inches(5.9), [
    ["model 的欄位", "說明"],
    ["Qwen/Qwen3-8B-AWQ", "Hugging Face 模型 ID"],
    ["qwen3-8b", "opencode 裡顯示的名字（英數字）"],
    ["32768", "上下文；opencode 至少 16384"],
    ["hermes", "tool parser，依模型系列選（下一頁）"],
], [2.3, 3.4], size=13, rowh=0.48, mono_cols=(0,))

# =============================================================================
#  20b. 使用 Hugging Face 上的模型
# =============================================================================
s = content_slide("06 管理與排錯 · 換模型", "使用 Hugging Face 上的模型（例：Llama 3.2）",
                  "vLLM 直接從 Hugging Face 下載模型，所以大部分模型都能用，但要符合幾個條件：要用 Instruct 版；"
                  "授權模型（例如 Meta 的 Llama）要先在模型頁面同意條款並設定 token；格式用原始權重或 AWQ / FP8，GGUF 不建議；"
                  "顯存要放得下；而且要支援工具呼叫並選對 parser。小模型通常比預設的 Qwen3-4B 弱，一定要跑 bench 比較。")
table(s, MX, TOP, Inches(6.9), [
    ["條件", "以 Llama 3.2 為例"],
    ["要用 Instruct 版", "Llama-3.2-3B-Instruct"],
    ["授權模型先同意條款 + token", "Meta Llama 要；Qwen 不用"],
    ["顯存放得下（12GB 約到 8B 量化）", "1B、3B 可以；11B Vision 不行"],
    ["格式：原始權重 / AWQ / FP8", "GGUF 不建議"],
    ["支援 tool calling + 選對 parser", "llama3_json"],
], [3.5, 3.4], size=13, rowh=0.5)
code(s, MX, TOP + Inches(3.25), Inches(6.9), Inches(1.35), [
    ">set HF_TOKEN hf_xxxxxxxx",
    ">model meta-llama/Llama-3.2-3B-Instruct llama3.2-3b 32768 llama3_json",
    ">bench",
], size=12)
table(s, MX + Inches(7.2), TOP, CW - Inches(7.2), [
    ["模型系列", "tool parser"],
    ["Qwen2.5 / Qwen3", "hermes"],
    ["Qwen3-Coder", "qwen3_coder"],
    ["Llama 3.1 / 3.2 / 3.3", "llama3_json"],
    ["Llama 4", "llama4_pythonic"],
    ["Mistral", "mistral"],
    ["Gemma 4", "gemma4"],
    ["GLM-4.5", "glm45"],
    ["gpt-oss", "openai"],
], [2.6, 2.3], size=13, rowh=0.46, mono_cols=(1,))
text(s, MX + Inches(7.2), TOP + Inches(4.25), CW - Inches(7.2), Inches(0.4),
     "選錯 parser：回應會直接出現 <tool_call> 文字", size=12, color=MUTED)

# =============================================================================
#  20c. 換模型的注意事項
# =============================================================================
s = content_slide("06 管理與排錯 · 換模型", "換模型的注意事項",
                  "切換模型時 vLLM 會重啟，所有人都會暫停約一分鐘，挑沒人用的時間或先通知。"
                  "下載過的模型會一直留在 WSL 裡佔硬碟，forget 不會刪權重。啟動失敗多半是顯存不足，調小上下文或換小模型，"
                  "switch 隨時可以換回原本能用的模型。repo 更新後要先跑 setup-host -SkipVllm 同步管理腳本，"
                  "加過 VPN 網段的話記得帶上同樣的 -AllowFrom。")
notes_list = [("⏸️", "切換時中斷約 1 分鐘", "所有人（opencode、Discord）都會暫停；挑沒人用的時間或先通知"),
              ("💾", "模型會佔硬碟", "forget 不刪權重；用 hf cache ls 查、hf cache rm 刪（刪資料夾不會釋放）"),
              ("🧯", "啟動失敗多半是顯存不足", "set SERVE_MAX_LEN 16384 或換小模型；switch 隨時換回能用的"),
              ("🏷️", "名稱改了要更新 client", "同學重跑 install-client；servers.json 的 models 也要改"),
              ("🔄", "先同步管理腳本", "repo 更新後執行 setup-host.cmd -SkipVllm（不會重裝 vLLM）"),
              ("🧱", "別讓防火牆被重設", "加過 VPN 網段：同步時也帶上 -AllowFrom")]
cw, ch = (CW - Inches(0.3)) / 2, Inches(1.35)
for i, (ic, t, d) in enumerate(notes_list):
    x, y = MX + (cw + Inches(0.3)) * (i % 2), TOP + (ch + Inches(0.2)) * (i // 2)
    box(s, x, y, cw, ch)
    badge(s, x + Inches(0.25), y + Inches(0.32), Inches(0.65), ic, fill=TINT, size=18)
    text(s, x + Inches(1.1), y + Inches(0.22), cw - Inches(1.3), Inches(1.0),
         [[(t, {"bold": True, "color": INK, "size": 16})], [(d, {"size": 13})]], space=4, line=1.1)

# =============================================================================
#  21. 安全性
# =============================================================================
s = content_slide("06 管理與排錯 · 安全性", "四層保護",
                  "校內使用時，API 金鑰加防火牆網段限制已經足夠。開 Funnel 時，全世界都連得到網址，"
                  "所以只公開 /v1、用 HTTPS，並且給校外的人之前先換金鑰，用完就關。")
layers = [("🔑", "API 金鑰", "每個請求都要帶 Authorization: Bearer 金鑰；外流就 new-key。"),
          ("🧱", "防火牆網段", "8000 只開給實驗室網段 + 學校 VPN 網段，不用 Any。"),
          ("🛡️", "Funnel 只公開 /v1", "vLLM 金鑰只保護 /v1；HTTPS 由 Tailscale 自動處理。"),
          ("🔒", "資料留在實驗室", "模型、程式碼、對話都不經過雲端 API。")]
cw, ch = (CW - Inches(0.3)) / 2, Inches(2.0)
for i, (ic, t, d) in enumerate(layers):
    x, y = MX + (cw + Inches(0.3)) * (i % 2), TOP + (ch + Inches(0.3)) * (i // 2)
    box(s, x, y, cw, ch)
    badge(s, x + Inches(0.35), y + Inches(0.35), Inches(0.7), ic, fill=TINT, size=20)
    text(s, x + Inches(1.35), y + Inches(0.35), cw - Inches(1.7), Inches(1.4),
         [[(t, {"bold": True, "color": INK, "size": 19})], [(d, {"size": 15})]], space=6, line=1.1)

# =============================================================================
#  22. 疑難排解
# =============================================================================
s = content_slide("06 管理與排錯 · 任何問題先看 vllmctl.cmd → logs", "疑難排解",
                  "最常見的是顯存不足和連不到。紀錄裡的 Engine core initialization failed 只是總結，"
                  "往上找第一個 RuntimeError 才是真正原因。")
table(s, MX, TOP, CW, [
    ["症狀", "處理"],
    ["opencode：Cannot connect to API", "vllmctl.cmd → status；沒在跑就 start"],
    ["重開機後大家都連不到", "要先登入 Windows；排程 vLLM-WSL-KeepAlive 要在執行中"],
    ["同學連不到、筆電本機可以", "IP 變了？筆電接手機熱點？在家要先連 VPN、網段加進防火牆"],
    ["啟動失敗、顯存不足（OOM）", "vllmctl.cmd → set SERVE_MAX_LEN 16384 或 set SERVE_GPU_UTIL 0.80"],
    ["opencode 報超過 context 上限", "/compact 或 /new；常發生就加大 SERVE_MAX_LEN"],
    ["Discord：連不上 vLLM / 金鑰錯誤", "vllmctl.cmd → status；.env 的 LLM_API_KEY 清空後重啟 bot"],
    ["桌面版 port already in use", "主機筆電才會遇到：重跑 install-client 會自動修正"],
], [3.8, 6.2], size=14, rowh=0.52)

# =============================================================================
#  23. 解除安裝
# =============================================================================
s = content_slide("06 管理與排錯 · 解除安裝", "解除安裝並還原原本的設定",
                  "setup-host 第一次安裝時，會先把電源、WSL 網路模式、OpenSSH 的原本狀態備份起來。"
                  "uninstall-host 就照這份備份還原，並移除 vLLM、排程、防火牆規則和這台的 opencode。"
                  "用舊版安裝、沒有備份的話，電源會改回 Windows 預設值，結束時會列出來。repo 的檔案和你自己下載的模型都不會動。")
code(s, MX, TOP, Inches(5.6), Inches(4.4), [
    "# 筆電（主機），會要求系統管理員權限",
    ">server\\wsl2\\uninstall-host.cmd",
    "# 看過清單後輸入 y",
    "",
    "# 常用參數",
    ">-KeepModels      # 保留模型",
    ">-KeepOpencode    # 保留這台的 opencode",
    ">-RemoveWsl       # 整個刪除 Ubuntu",
    "",
    "# 同學的電腦",
    ">opencode\\client\\uninstall-client.cmd",
], size=13)
table(s, MX + Inches(5.9), TOP, CW - Inches(5.9), [
    ["項目", "移除 / 還原"],
    ["排程", "vLLM 常駐、Discord bot"],
    ["WSL 裡的 vLLM", "服務、設定、金鑰、下載的模型"],
    ["防火牆", "vLLM-opencode-8000"],
    [".wslconfig", "網路模式還原成安裝前"],
    ["電源", "睡眠 / 休眠 / 闔上螢幕還原"],
    ["Funnel · OpenSSH", "關閉 / 還原成安裝前"],
    ["opencode", "程式與實驗室設定（其他設定保留）"],
], [2.2, 3.6], size=13, rowh=0.5)
text(s, MX + Inches(5.9), TOP + Inches(4.15), CW - Inches(5.9), Inches(0.4),
     "備份位置：%LOCALAPPDATA%\\vllm-opencode-backup", size=12, color=MUTED)

# =============================================================================
#  24. 指令小抄
# =============================================================================
s = new_slide(dark=True)
header(s, "總結", "指令小抄", dark=True)
footer(s, dark=True)
notes(s, "最後這頁可以截圖貼在實驗室群組。完整步驟寫在 repo 根目錄的「完整步驟.md」。")
cols = [("安裝", [("筆電主機（server\\wsl2）", "setup-host.cmd"),
                  ("同學電腦（opencode\\client）", "install-client.cmd"),
                  ("Discord bot（discord\\client）", "deploy.ps1 → start-bot.bat")]),
        ("使用", [("在專案資料夾", "opencode"),
                  ("opencode 裡", "/init  /models  /compact  /undo"),
                  ("Discord", "!ppt  !word  !excel  !web")]),
        ("管理（vllmctl.cmd）", [("連線資訊 / 紀錄", "info · logs"),
                                ("換模型", "preset quality · bench"),
                                ("換金鑰 / 校外", "new-key · funnel.cmd on|off")])]
cw = (CW - Inches(0.5)) / 3
for i, (t, rows) in enumerate(cols):
    x = MX + (cw + Inches(0.25)) * i
    box(s, x, TOP, cw, Inches(4.6), fill=NAVY2, line=None)
    text(s, x + Inches(0.35), TOP + Inches(0.3), cw - Inches(0.7), Inches(0.5), t, size=20, color=ORANGE, bold=True)
    for j, (lab, cmd) in enumerate(rows):
        y = TOP + Inches(1.05) + Inches(1.15) * j
        text(s, x + Inches(0.35), y, cw - Inches(0.7), Inches(1.0),
             [[(lab, {"size": 13, "color": MUTED})], [(cmd, {"size": 14, "color": WHITE, "mono": True})]], space=4)

prs.save(OUT)
print(OUT, len(prs.slides), "頁")
