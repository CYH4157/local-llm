"""
office_gen.py — 把模型輸出的 JSON 排版成 .pptx / .docx / .xlsx。

流程：模型只負責「內容」（受 JSON schema 約束），排版全部交給 Python，
所以產出格式穩定。若 templates/ 內有範本檔就套用範本（校徽、配色、頁首頁尾）：
    templates/ppt_template.pptx   — 第 0 個版面=標題投影片，第 1 個版面=標題及內容
    templates/word_template.docx  — 會沿用其樣式、頁首頁尾，內文清空後重填
"""

import os
import re
from datetime import datetime

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn as docx_qn
from docx.shared import Cm, Pt, RGBColor as DocxRGB
from lxml import etree
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn as pptx_qn
from pptx.util import Emu, Inches, Pt as PPt

BASE_DIR     = os.path.dirname(os.path.abspath(__file__))
TEMPLATE_DIR = os.path.join(BASE_DIR, "templates")
OUTPUT_DIR   = os.path.join(BASE_DIR, "outputs")

FONT_EA    = os.getenv("DOC_FONT_EA", "標楷體")          # Word 中文字型（國科會格式常用）
FONT_LATIN = os.getenv("DOC_FONT_LATIN", "Times New Roman")
PPT_FONT   = os.getenv("PPT_FONT", "微軟正黑體")

NAVY   = RGBColor(0x1F, 0x3A, 0x5F)
ACCENT = RGBColor(0x2E, 0x86, 0xAB)
LIGHT  = RGBColor(0xF4, 0xF6, 0xF8)
GREY   = RGBColor(0x6B, 0x72, 0x80)
WHITE  = RGBColor(0xFF, 0xFF, 0xFF)


# =============================================================================
#  JSON schema（給 vLLM guided decoding 用）
# =============================================================================
_STR_LIST = {"type": "array", "items": {"type": "string"}}
_TABLE = {
    "type": "object",
    "properties": {
        "headers": _STR_LIST,
        "rows": {"type": "array", "items": _STR_LIST},
    },
    "required": ["headers", "rows"],
}

PPT_SCHEMA = {
    "type": "object",
    "properties": {
        "title": {"type": "string"},
        "subtitle": {"type": "string"},
        "slides": {
            "type": "array",
            "minItems": 3,
            "maxItems": 20,
            "items": {
                "type": "object",
                "properties": {
                    "title": {"type": "string"},
                    "bullets": {"type": "array", "items": {"type": "string"}, "maxItems": 7},
                    "notes": {"type": "string"},
                },
                "required": ["title", "bullets", "notes"],
            },
        },
    },
    "required": ["title", "subtitle", "slides"],
}

DOCX_SCHEMA = {
    "type": "object",
    "properties": {
        "title": {"type": "string"},
        "subtitle": {"type": "string"},
        "sections": {
            "type": "array",
            "minItems": 1,
            "items": {
                "type": "object",
                "properties": {
                    "heading": {"type": "string"},
                    "paragraphs": _STR_LIST,
                    "bullets": _STR_LIST,
                    "table": _TABLE,
                },
                "required": ["heading", "paragraphs", "bullets", "table"],
            },
        },
        "references": _STR_LIST,
    },
    "required": ["title", "subtitle", "sections", "references"],
}

XLSX_SCHEMA = {
    "type": "object",
    "properties": {
        "title": {"type": "string"},
        "sheets": {
            "type": "array",
            "minItems": 1,
            "maxItems": 5,
            "items": {
                "type": "object",
                "properties": {
                    "name": {"type": "string"},
                    "headers": _STR_LIST,
                    "rows": {
                        "type": "array",
                        "items": {"type": "array",
                                  "items": {"anyOf": [{"type": "string"}, {"type": "number"}]}},
                    },
                    "notes": {"type": "string"},
                },
                "required": ["name", "headers", "rows", "notes"],
            },
        },
    },
    "required": ["title", "sheets"],
}


# =============================================================================
#  提示詞
# =============================================================================
_COMMON = (
    "你是大學研究室的文件助理。只輸出 JSON，不要任何說明文字。"
    "一律使用台灣繁體中文與台灣慣用術語（專有名詞可保留英文）。"
)

_WORD_MODES = {
    "計畫": (
        "撰寫一份『國科會專題研究計畫』草稿。sections 依序包含："
        "計畫摘要、研究背景與動機、國內外相關研究（文獻探討）、研究目的、"
        "研究方法與步驟、預期成果與貢獻、執行進度（table 用甘特表：headers 為工作項目與各月份，格子填 ■ 或空字串）。"
        "每節 paragraphs 2~4 段、每段 120~250 字，論述要具體、有邏輯。"
    ),
    "報告": (
        "撰寫一份大學/研究所『課堂報告』。sections 依序包含："
        "摘要、前言、主體內容（可拆成 2~4 節，標題自訂）、討論、結論。"
        "每節 paragraphs 1~3 段、每段 100~200 字。"
    ),
    "一般": (
        "依使用者需求撰寫文件，自行規劃 3~8 個章節。"
        "每節 paragraphs 1~3 段，需要列點時放 bullets，需要表格時放 table。"
    ),
}

_REF_RULE = (
    "references 規則：你無法查證文獻，禁止捏造作者、年份或期刊。"
    "若使用者提供的參考資料中有文獻就照抄；否則每筆寫成「[待補] 建議查詢：<關鍵字或主題>」。"
    "不需要表格的章節，table 的 headers 與 rows 都給空陣列；不需要列點就給空陣列。"
)


def _user_block(request: str, ref_text: str) -> str:
    s = f"需求：{request}"
    if ref_text:
        s += f"\n\n以下是使用者提供的參考資料，內容請以此為準：\n---\n{ref_text}\n---"
    return s


def ppt_messages(request: str, ref_text: str = "") -> list[dict]:
    sys = (
        _COMMON
        + "請規劃一份簡報：title 為簡報標題、subtitle 為副標（如課程名稱或報告人，不知道就寫主題簡述）。"
        "slides 為內容頁（不含封面），通常 6~12 頁，最後一頁為結論或總結。"
        "每頁 bullets 3~6 點、每點 10~40 字，精簡有力，不要整段文章。"
        "notes 為該頁講者備註（口語化的講稿，80~200 字）。"
    )
    return [{"role": "system", "content": sys},
            {"role": "user", "content": _user_block(request, ref_text)}]


def docx_messages(mode: str, request: str, ref_text: str = "") -> list[dict]:
    sys = _COMMON + _WORD_MODES.get(mode, _WORD_MODES["一般"]) + _REF_RULE
    return [{"role": "system", "content": sys},
            {"role": "user", "content": _user_block(request, ref_text)}]


def xlsx_messages(request: str, ref_text: str = "") -> list[dict]:
    sys = (
        _COMMON
        + "請設計 Excel 表格。每個 sheet：headers 是第 1 列欄位名稱，rows 從第 2 列開始。"
        "數值請直接用 number，不要加單位或千分位（單位寫在欄位名稱，如「金額(元)」）。"
        "需要計算時可用 Excel 公式字串，例如 \"=SUM(C2:C6)\"、\"=B2*C2\"，列號要對應正確。"
        "notes 寫這張表的使用說明（沒有就空字串）。"
    )
    return [{"role": "system", "content": sys},
            {"role": "user", "content": _user_block(request, ref_text)}]


# =============================================================================
#  共用工具
# =============================================================================
def output_path(title: str, ext: str) -> str:
    os.makedirs(OUTPUT_DIR, exist_ok=True)
    safe = re.sub(r'[\\/:*?"<>|\r\n\t]+', "_", title or "document").strip(" ._")[:40] or "document"
    stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    return os.path.join(OUTPUT_DIR, f"{safe}_{stamp}.{ext}")


def _template(name: str) -> str | None:
    p = os.path.join(TEMPLATE_DIR, name)
    return p if os.path.isfile(p) else None


def _clean_list(items) -> list[str]:
    return [str(x).strip() for x in (items or []) if str(x).strip()]


# =============================================================================
#  PowerPoint
# =============================================================================
def _ppt_font(run, size=None, bold=None, color=None):
    if color is not None:
        run.font.color.rgb = color
    if size is not None:
        run.font.size = PPt(size)
    if bold is not None:
        run.font.bold = bold
    run.font.name = PPT_FONT
    # 中文字型要另外寫在 <a:ea>，只設 font.name 對中文無效
    rPr = run._r.get_or_add_rPr()
    for old in rPr.findall(pptx_qn("a:ea")):
        rPr.remove(old)
    ea = etree.SubElement(rPr, pptx_qn("a:ea"))
    ea.set("typeface", PPT_FONT)
    latin = rPr.find(pptx_qn("a:latin"))
    if latin is not None:
        latin.addnext(ea)


def _bullet(paragraph):
    pPr = paragraph._p.get_or_add_pPr()
    pPr.set("marL", "342900")
    pPr.set("indent", "-342900")
    bu = etree.SubElement(pPr, pptx_qn("a:buChar"))
    bu.set("char", "•")


def _textbox(slide, left, top, width, height):
    tb = slide.shapes.add_textbox(left, top, width, height)
    tf = tb.text_frame
    tf.word_wrap = True
    return tf


def _rect(slide, left, top, width, height, color):
    shp = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, width, height)
    shp.fill.solid()
    shp.fill.fore_color.rgb = color
    shp.line.fill.background()
    shp.shadow.inherit = False   # 拿掉預設陰影
    return shp


def _bullet_size(bullets: list[str]) -> int:
    n, chars = len(bullets), sum(len(b) for b in bullets)
    if n <= 4 and chars <= 140:
        return 24
    if n <= 6 and chars <= 240:
        return 20
    return 16


def _drop_all_slides(prs):
    sldIdLst = prs.slides._sldIdLst
    for sldId in list(sldIdLst):
        prs.part.drop_rel(sldId.rId)
        sldIdLst.remove(sldId)


def _set_notes(slide, notes: str):
    if notes and notes.strip():
        slide.notes_slide.notes_text_frame.text = notes.strip()


def build_pptx(data: dict) -> str:
    title = data.get("title", "簡報").strip() or "簡報"
    subtitle = data.get("subtitle", "").strip()
    slides = data.get("slides", [])
    path = output_path(title, "pptx")

    tpl = _template("ppt_template.pptx")
    if tpl:
        _build_pptx_from_template(tpl, title, subtitle, slides, path)
    else:
        _build_pptx_default(title, subtitle, slides, path)
    return path


def _build_pptx_from_template(tpl, title, subtitle, slides, path):
    prs = Presentation(tpl)
    _drop_all_slides(prs)
    s = prs.slides.add_slide(prs.slide_layouts[0])
    if s.shapes.title is not None:
        s.shapes.title.text = title
    for ph in s.placeholders:
        if ph.placeholder_format.idx == 1:
            ph.text = subtitle
    content_layout = prs.slide_layouts[1 if len(prs.slide_layouts) > 1 else 0]
    for sd in slides:
        s = prs.slides.add_slide(content_layout)
        if s.shapes.title is not None:
            s.shapes.title.text = sd.get("title", "")
        body = next((ph for ph in s.placeholders if ph.placeholder_format.idx == 1), None)
        bullets = _clean_list(sd.get("bullets"))
        if body is not None:
            tf = body.text_frame
            tf.text = bullets[0] if bullets else ""
            for b in bullets[1:]:
                tf.add_paragraph().text = b
        _set_notes(s, sd.get("notes", ""))
    prs.save(path)


def _build_pptx_default(title, subtitle, slides, path):
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    W, H = prs.slide_width, prs.slide_height
    blank = prs.slide_layouts[6]
    margin = Inches(0.7)

    # --- 封面 ---
    s = prs.slides.add_slide(blank)
    _rect(s, 0, 0, W, H, NAVY)
    _rect(s, margin, Inches(4.35), Inches(1.6), Emu(54000), ACCENT)
    tf = _textbox(s, margin, Inches(2.2), W - margin * 2, Inches(2.0))
    tf.vertical_anchor = MSO_ANCHOR.BOTTOM
    r = tf.paragraphs[0].add_run()
    r.text = title
    _ppt_font(r, 40, True, WHITE)
    tf = _textbox(s, margin, Inches(4.6), W - margin * 2, Inches(1.2))
    r = tf.paragraphs[0].add_run()
    r.text = subtitle
    _ppt_font(r, 20, False, RGBColor(0xD6, 0xE4, 0xF0))
    p = tf.add_paragraph()
    r = p.add_run()
    r.text = datetime.now().strftime("%Y / %m / %d")
    _ppt_font(r, 14, False, RGBColor(0xA9, 0xBC, 0xD0))

    # --- 內容頁 ---
    total = len(slides)
    for i, sd in enumerate(slides, 1):
        s = prs.slides.add_slide(blank)
        _rect(s, 0, 0, W, Inches(1.25), LIGHT)
        _rect(s, 0, Inches(1.25), W, Emu(38000), ACCENT)
        _rect(s, 0, 0, Emu(110000), Inches(1.25), NAVY)

        tf = _textbox(s, margin, Inches(0.2), W - margin * 2, Inches(0.9))
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        r = tf.paragraphs[0].add_run()
        r.text = sd.get("title", "")
        _ppt_font(r, 30, True, NAVY)

        bullets = _clean_list(sd.get("bullets"))
        size = _bullet_size(bullets)
        tf = _textbox(s, margin, Inches(1.65), W - margin * 2, H - Inches(2.4))
        for j, b in enumerate(bullets):
            p = tf.paragraphs[0] if j == 0 else tf.add_paragraph()
            p.space_after = PPt(size * 0.6)
            _bullet(p)
            r = p.add_run()
            r.text = b
            _ppt_font(r, size, False, RGBColor(0x22, 0x22, 0x22))

        tf = _textbox(s, W - Inches(1.6), H - Inches(0.55), Inches(1.2), Inches(0.4))
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.RIGHT
        r = p.add_run()
        r.text = f"{i} / {total}"
        _ppt_font(r, 11, False, GREY)

        _set_notes(s, sd.get("notes", ""))

    prs.save(path)


# =============================================================================
#  Word
# =============================================================================
def _set_style_font(style, size=None, bold=None, color=None):
    style.font.name = FONT_LATIN
    if size is not None:
        style.font.size = Pt(size)
    if bold is not None:
        style.font.bold = bold
    if color is not None:
        style.font.color.rgb = color
    rPr = style.element.get_or_add_rPr()
    rFonts = rPr.find(docx_qn("w:rFonts"))
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts")
        rPr.insert(0, rFonts)
    rFonts.set(docx_qn("w:eastAsia"), FONT_EA)
    # 標題樣式預設用佈景主題字型，要清掉才會套用我們指定的字型
    for attr in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
        rFonts.attrib.pop(docx_qn(attr), None)


def _add_page_number(doc):
    p = doc.sections[0].footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    for tag, text in (("begin", None), (None, "PAGE"), ("end", None)):
        if tag:
            el = OxmlElement("w:fldChar")
            el.set(docx_qn("w:fldCharType"), tag)
        else:
            el = OxmlElement("w:instrText")
            el.set(docx_qn("xml:space"), "preserve")
            el.text = text
        run._r.append(el)


def _style_or(doc, name, fallback="Normal"):
    try:
        return doc.styles[name]
    except KeyError:
        return doc.styles[fallback]


def _shade(cell, hex_color):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(docx_qn("w:val"), "clear")
    shd.set(docx_qn("w:color"), "auto")
    shd.set(docx_qn("w:fill"), hex_color)
    tcPr.append(shd)


def _docx_table(doc, table: dict):
    headers = _clean_list(table.get("headers"))
    rows = [[str(c) for c in r] for r in (table.get("rows") or []) if r]
    if not headers:
        return
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = _style_or(doc, "Table Grid")
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for c, h in zip(t.rows[0].cells, headers):
        c.text = h
        c.paragraphs[0].runs[0].bold = True
        c.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        _shade(c, "D9E2F3")
    for r in rows:
        cells = t.add_row().cells
        for c, v in zip(cells, r):
            c.text = v
            if len(v) <= 3:   # ■、數字等短內容置中
                c.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_paragraph()


def build_docx(data: dict) -> str:
    title = data.get("title", "文件").strip() or "文件"
    path = output_path(title, "docx")

    tpl = _template("word_template.docx")
    if tpl:
        doc = Document(tpl)
        body = doc.element.body
        for child in list(body):
            if child.tag != docx_qn("w:sectPr"):
                body.remove(child)
    else:
        doc = Document()
        for sec in doc.sections:
            sec.page_height, sec.page_width = Cm(29.7), Cm(21.0)   # A4
            sec.left_margin = sec.right_margin = Cm(2.54)
            sec.top_margin = sec.bottom_margin = Cm(2.54)
        _set_style_font(doc.styles["Normal"], 12)
        doc.styles["Normal"].paragraph_format.line_spacing = 1.5
        black = DocxRGB(0, 0, 0)
        _set_style_font(doc.styles["Title"], 20, True, black)
        for lvl, size in ((1, 16), (2, 14), (3, 12)):
            _set_style_font(doc.styles[f"Heading {lvl}"], size, True, black)
        for name in ("List Bullet", "List Number"):
            _set_style_font(doc.styles[name], 12)
        _add_page_number(doc)

    t = doc.add_paragraph(title, style=_style_or(doc, "Title"))
    t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle = data.get("subtitle", "").strip()
    if subtitle:
        p = doc.add_paragraph(subtitle)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER

    for sec in data.get("sections", []):
        doc.add_paragraph(sec.get("heading", ""), style=_style_or(doc, "Heading 1"))
        for para in _clean_list(sec.get("paragraphs")):
            p = doc.add_paragraph(para)
            p.paragraph_format.first_line_indent = Pt(24)   # 首行縮排 2 字
        for b in _clean_list(sec.get("bullets")):
            doc.add_paragraph(b, style=_style_or(doc, "List Bullet"))
        if sec.get("table"):
            _docx_table(doc, sec["table"])

    refs = _clean_list(data.get("references"))
    if refs:
        doc.add_paragraph("參考文獻", style=_style_or(doc, "Heading 1"))
        for i, ref in enumerate(refs, 1):
            doc.add_paragraph(f"[{i}] {ref}")

    doc.save(path)
    return path


# =============================================================================
#  Excel
# =============================================================================
_NUM_RE = re.compile(r"^-?\d+(\.\d+)?$")


def _cell_value(v):
    if isinstance(v, str):
        s = v.strip().replace(",", "")
        if _NUM_RE.match(s):
            return float(s) if "." in s else int(s)
        return v.strip()
    return v


def _display_width(v) -> int:
    s = str(v)
    return sum(2 if ord(ch) > 0x2E80 else 1 for ch in s)


def _sheet_name(name: str, used: set) -> str:
    base = re.sub(r"[\[\]:*?/\\]", "_", name or "Sheet").strip("'")[:31] or "Sheet"
    cand, n = base, 2
    while cand in used:
        suffix = f"_{n}"
        cand, n = base[:31 - len(suffix)] + suffix, n + 1
    used.add(cand)
    return cand


def build_xlsx(data: dict) -> str:
    title = data.get("title", "表格").strip() or "表格"
    path = output_path(title, "xlsx")

    wb = Workbook()
    wb.remove(wb.active)
    used: set[str] = set()
    thin = Side(style="thin", color="B7C3D0")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    head_fill = PatternFill("solid", fgColor="1F3A5F")

    for sh in data.get("sheets", []):
        ws = wb.create_sheet(_sheet_name(sh.get("name", ""), used))
        headers = [str(h) for h in sh.get("headers", [])]
        rows = [[_cell_value(v) for v in r] for r in sh.get("rows", [])]
        ncols = max([len(headers)] + [len(r) for r in rows] + [1])

        ws.append(headers)
        for r in rows:
            ws.append(r)

        for c in range(1, len(headers) + 1):
            cell = ws.cell(row=1, column=c)
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = head_fill
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        for row in ws.iter_rows(min_row=1, max_row=len(rows) + 1, max_col=ncols):
            for cell in row:
                cell.border = border
                if cell.row > 1:
                    cell.alignment = Alignment(vertical="center", wrap_text=True)

        for c in range(1, ncols + 1):
            col = [headers[c - 1] if c <= len(headers) else ""] + \
                  [r[c - 1] for r in rows if c <= len(r) and not str(r[c - 1]).startswith("=")]
            width = max((_display_width(v) for v in col), default=8)
            ws.column_dimensions[get_column_letter(c)].width = min(max(width + 2, 8), 50)

        ws.freeze_panes = "A2"
        if headers and rows:
            ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{len(rows) + 1}"
        notes = (sh.get("notes") or "").strip()
        if notes:
            cell = ws.cell(row=len(rows) + 3, column=1, value=f"說明：{notes}")
            cell.font = Font(italic=True, color="6B7280")

    if not wb.sheetnames:
        wb.create_sheet("Sheet1")
    wb.save(path)
    return path
