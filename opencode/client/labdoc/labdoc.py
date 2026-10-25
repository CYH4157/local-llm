#!/usr/bin/env python3
"""
labdoc — 實驗室 opencode 的文件小工具（由 install-client 安裝，opencode 的 /指令 會呼叫它）

模型只負責寫 Markdown（內容），排版、轉檔全部交給這支程式，小模型也能穩定產出 Office 檔。

    labdoc read     檔案                 讀 PDF / Word / PPT / Excel / CSV / 文字檔，印出純文字
    labdoc pdf2word 檔案.pdf             PDF 轉 Word（保留版面）
    labdoc word     檔案.md              Markdown → Word
    labdoc slides   檔案.md              Markdown → PowerPoint
    labdoc excel    檔案.md|.csv         Markdown 表格 / CSV → Excel（--total 加合計列）
    labdoc topdf    檔案.docx|.pptx|.xlsx  Office → PDF（需要 Office 或 LibreOffice）
    labdoc pdf merge|pages|delete|rotate|fromimg|topng|compress ...
    labdoc tw       檔案                 檢查簡體字、大陸用語、半形標點（--fix 另存修正版）

產生的檔案預設放在輸入檔旁邊；同名檔案已存在時自動加 (1)、(2)，不會覆蓋（--force 才覆蓋）。
"""
import argparse
import csv
import io
import os
import re
import shutil
import subprocess
import sys
from pathlib import Path

# opencode 讀的是 UTF-8；Windows 預設是 cp950，中文會變亂碼或直接出錯
for _s in (sys.stdout, sys.stderr):
    try:
        _s.reconfigure(encoding="utf-8", errors="replace")
    except Exception:
        pass

FONT_EA    = os.getenv("LABDOC_FONT_EA", "標楷體")           # Word 中文字型（公文常用）
FONT_LATIN = os.getenv("LABDOC_FONT_LATIN", "Times New Roman")
PPT_FONT   = os.getenv("LABDOC_PPT_FONT", "微軟正黑體")
MAX_CHARS  = 15000   # read 預設最多印出的字數（本地模型上下文約 32k token）


class UserError(Exception):
    """給使用者看的錯誤：只印訊息，不印 traceback"""


# =============================================================================
#  共用
# =============================================================================
def need(path):
    p = Path(path).expanduser()
    if not p.exists():
        raise UserError(f"找不到檔案：{path}（請確認檔名，或先 cd 到檔案所在的資料夾）")
    return p


def out_path(src, ext, out=None, force=False, suffix=""):
    """決定輸出檔名；不覆蓋既有檔案（除非 --force）"""
    p = Path(out).expanduser() if out else Path(src).with_name(Path(src).stem + suffix + ext)
    if p.suffix.lower() != ext:
        p = p.with_suffix(ext)
    if force or not p.exists():
        return p
    n = 1
    while True:
        q = p.with_name(f"{p.stem} ({n}){p.suffix}")
        if not q.exists():
            return q
        n += 1


def done(p):
    print(f"已產生：{Path(p).resolve()}")


def page_list(spec, total):
    """'1-3,5,8-' → [0,1,2,4,7,...]（0 起算）；None = 全部"""
    if not spec:
        return list(range(total))
    pages = []
    for part in str(spec).replace("，", ",").split(","):
        part = part.strip()
        if not part:
            continue
        m = re.fullmatch(r"(\d*)\s*[-~]\s*(\d*)", part)
        if m:
            a = int(m.group(1)) if m.group(1) else 1
            b = int(m.group(2)) if m.group(2) else total
        elif part.isdigit():
            a = b = int(part)
        else:
            raise UserError(f"看不懂頁碼「{part}」，請用 1-3,5 這種格式")
        if a < 1 or b > total or a > b:
            raise UserError(f"頁碼 {part} 超出範圍（這份共 {total} 頁）")
        pages.extend(range(a - 1, b))
    return pages


def read_text_file(p):
    raw = p.read_bytes()
    for enc in ("utf-8-sig", "cp950", "big5hkscs", "gb18030"):
        try:
            return raw.decode(enc)
        except UnicodeDecodeError:
            continue
    return raw.decode("utf-8", errors="replace")


def md_table(rows):
    rows = [[("" if c is None else str(c)).replace("\n", " ").replace("|", "／") for c in r] for r in rows]
    rows = [r for r in rows if any(c.strip() for c in r)]
    if not rows:
        return ""
    w = max(len(r) for r in rows)
    rows = [r + [""] * (w - len(r)) for r in rows]
    lines = ["| " + " | ".join(rows[0]) + " |", "|" + "---|" * w]
    lines += ["| " + " | ".join(r) + " |" for r in rows[1:]]
    return "\n".join(lines)


# =============================================================================
#  read：各種檔案 → 純文字
# =============================================================================
def read_pdf(p, pages_spec):
    import pymupdf as fitz
    doc = fitz.open(p)
    pages = page_list(pages_spec, doc.page_count)
    out, empty = [f"（PDF 共 {doc.page_count} 頁）"], 0
    for i in pages:
        t = doc[i].get_text("text").strip()
        if len(t) < 20:
            empty += 1
        out.append(f"\n--- 第 {i + 1} 頁 ---\n{t}")
    if pages and empty == len(pages):
        out.insert(1, "【注意】這份 PDF 幾乎讀不到文字，可能是掃描檔或圖片（labdoc 不做 OCR）。"
                      "可以用 Word 開啟 PDF，或手機的文字辨識功能先轉成文字。")
    return "\n".join(out)


def read_docx(p):
    from docx import Document
    from docx.table import Table
    from docx.text.paragraph import Paragraph
    doc = Document(p)
    out = []
    for el in doc.element.body.iterchildren():
        tag = el.tag.split("}")[-1]
        if tag == "p":
            para = Paragraph(el, doc)
            t = para.text.strip()
            if not t:
                continue
            style = (para.style.name or "") if para.style is not None else ""
            m = re.match(r"(Heading|標題)\s*(\d)", style)
            if style == "Title":
                out.append(f"# {t}")
            elif m:
                out.append("#" * (int(m.group(2)) + 1) + " " + t)
            elif "List" in style or "清單" in style:
                out.append(f"- {t}")
            else:
                out.append(t)
        elif tag == "tbl":
            tbl = Table(el, doc)
            out.append(md_table([[c.text.strip() for c in r.cells] for r in tbl.rows]))
    return "\n\n".join(x for x in out if x)


def read_pptx(p):
    from pptx import Presentation
    prs = Presentation(p)
    out = [f"（共 {len(prs.slides)} 張投影片）"]
    for n, slide in enumerate(prs.slides, 1):
        out.append(f"\n## 第 {n} 張")
        for sh in slide.shapes:
            if sh.has_text_frame:
                for para in sh.text_frame.paragraphs:
                    t = "".join(r.text for r in para.runs).strip()
                    if t:
                        out.append("  " * para.level + "- " + t)
            if getattr(sh, "has_table", False) and sh.has_table:
                out.append(md_table([[c.text.strip() for c in r.cells] for r in sh.table.rows]))
        if slide.has_notes_slide:
            notes = slide.notes_slide.notes_text_frame.text.strip()
            if notes:
                out.append(f"> 備註：{notes}")
    return "\n".join(out)


def read_xlsx(p):
    from datetime import datetime, time
    from openpyxl import load_workbook
    wb = load_workbook(p, data_only=True, read_only=True)
    out = []
    for ws in wb.worksheets:
        rows = [[(c.date() if isinstance(c, datetime) and c.time() == time() else c) for c in r]
                for r in ws.iter_rows(values_only=True)]
        out.append(f"## 工作表：{ws.title}（{len(rows)} 列）\n\n" + md_table(rows))
    return "\n\n".join(out)


def cmd_read(a):
    p = need(a.file)
    ext = p.suffix.lower()
    if ext == ".pdf":
        text = read_pdf(p, a.pages)
    elif ext == ".docx":
        text = read_docx(p)
    elif ext == ".pptx":
        text = read_pptx(p)
    elif ext in (".xlsx", ".xlsm"):
        text = read_xlsx(p)
    elif ext == ".csv":
        text = md_table(list(csv.reader(io.StringIO(read_text_file(p)))))
    elif ext in (".doc", ".ppt", ".xls", ".odt", ".rtf"):
        raise UserError(f"{ext} 是舊格式，請先用 Office 另存成 .docx / .pptx / .xlsx，或執行 labdoc topdf 轉成 PDF 再讀")
    elif ext in (".png", ".jpg", ".jpeg", ".gif", ".bmp", ".heic"):
        raise UserError("這是圖片檔，labdoc 讀不到圖片裡的文字（不做 OCR）")
    else:
        text = read_text_file(p)
    if a.out:
        Path(a.out).write_text(text, encoding="utf-8")
        print(f"已把全文（{len(text)} 字）寫到：{Path(a.out).resolve()}")
        return
    if len(text) > a.max:
        print(text[:a.max])
        print(f"\n【只顯示前 {a.max} 字，全文共 {len(text)} 字】"
              + ("PDF 可用 --pages 指定頁碼分段讀，例如 --pages 6-10" if ext == ".pdf" else "可用 --max 調整"))
    else:
        print(text)


# =============================================================================
#  Markdown 解析（只支援常用語法：標題、段落、清單、表格、引言、圖片、分隔線、程式碼）
# =============================================================================
INLINE = re.compile(r"(\*\*.+?\*\*|__.+?__|`[^`]+`|\[[^\]]+\]\([^)]+\))")


def inline_runs(text):
    """'a **b** `c`' → [(text, bold, code)]"""
    runs = []
    for part in INLINE.split(text):
        if not part:
            continue
        if (part.startswith("**") and part.endswith("**")) or (part.startswith("__") and part.endswith("__")):
            runs.append((part[2:-2], True, False))
        elif part.startswith("`") and part.endswith("`"):
            runs.append((part[1:-1], False, True))
        elif part.startswith("["):
            m = re.match(r"\[([^\]]+)\]\(([^)]+)\)", part)
            runs.append((f"{m.group(1)}（{m.group(2)}）" if m.group(1) != m.group(2) else m.group(1), False, False))
        else:
            runs.append((part, False, False))
    return runs


def plain(text):
    return "".join(t for t, _, _ in inline_runs(text))


def is_cjk(ch):
    return "⺀" <= ch <= "鿿" or "＀" <= ch <= "￯" or "　" <= ch <= "〿"


def parse_md(text):
    blocks, para = [], []

    def flush():
        # 中文文件（公文、信件）習慣一行一段，所以每一行都是獨立段落
        for line in para:
            blocks.append({"k": "p", "text": line})
        para.clear()

    lines = text.replace("\t", "    ").splitlines()
    i = 0
    while i < len(lines):
        line = lines[i].rstrip()
        s = line.strip()
        if s.startswith("```"):
            flush()
            code = []
            i += 1
            while i < len(lines) and not lines[i].strip().startswith("```"):
                code.append(lines[i])
                i += 1
            blocks.append({"k": "code", "text": "\n".join(code)})
        elif not s:
            flush()
        elif re.match(r"^#{1,6}\s", s):
            flush()
            m = re.match(r"^(#{1,6})\s+(.*?)\s*#*$", s)
            blocks.append({"k": "h", "level": len(m.group(1)), "text": m.group(2)})
        elif re.fullmatch(r"(-\s*){3,}|(\*\s*){3,}|(_\s*){3,}", s):
            flush()
            blocks.append({"k": "hr"})
        elif s.startswith("|"):
            flush()
            rows = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                r = lines[i].strip().strip("|")
                if not re.fullmatch(r"[\s:|-]+", r):
                    rows.append([c.strip() for c in r.split("|")])
                i += 1
            i -= 1
            if rows:
                blocks.append({"k": "table", "rows": rows})
        elif re.match(r"^\s*([-*+]|\d+[.)、])\s+", line):
            flush()
            m = re.match(r"^(\s*)([-*+]|\d+[.)、])\s+(.*)", line)
            ordered = m.group(2)[0].isdigit()
            blocks.append({"k": "li", "level": min(len(m.group(1)) // 2, 3), "ordered": ordered,
                           "marker": m.group(2) if ordered else "", "text": m.group(3)})
        elif s.startswith(">"):
            flush()
            blocks.append({"k": "quote", "text": s.lstrip(">").strip()})
        elif re.fullmatch(r"!\[[^\]]*\]\([^)]+\)", s):
            flush()
            m = re.fullmatch(r"!\[([^\]]*)\]\(([^)\s]+)[^)]*\)", s)
            blocks.append({"k": "img", "alt": m.group(1), "path": m.group(2)})
        else:
            para.append(s)
        i += 1
    flush()
    return blocks


def to_number(s):
    """'NT$1,200' → 1200、'15%' → (0.15, '%')；不是數字回傳 None。電話、編號等前導 0 保留文字"""
    t = str(s).strip()
    if not t or re.fullmatch(r"0\d+", t):
        return None
    pct = t.endswith("%")
    t = re.sub(r"^(NT\$|NTD|\$|＄)\s*", "", t.rstrip("%").strip())
    t = re.sub(r"\s*(元|NTD|TWD)$", "", t).replace(",", "")
    if re.fullmatch(r"-?\d+", t):
        v = int(t)
    elif re.fullmatch(r"-?\d*\.\d+", t):
        v = float(t)
    else:
        return None
    return (v / 100, "%") if pct else (v, "")


def resolve_img(path, base):
    p = Path(path)
    if not p.is_absolute():
        p = Path(base) / p
    return p if p.exists() else None


# =============================================================================
#  Markdown → Word
# =============================================================================
def docx_font(el, latin, ea, size=None, bold=None, color=None):
    """設定 style 或 run 的中英文字型（清掉主題字型，否則 Word 會忽略）"""
    from docx.oxml.ns import qn
    rpr = el.get_or_add_rPr()
    rf = rpr.find(qn("w:rFonts"))
    if rf is None:
        rf = rpr.makeelement(qn("w:rFonts"), {})
        rpr.insert(0, rf)
    for att in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
        rf.attrib.pop(qn(att), None)
    rf.set(qn("w:ascii"), latin)
    rf.set(qn("w:hAnsi"), latin)
    rf.set(qn("w:eastAsia"), ea)


def add_runs(par, text, size=None):
    from docx.shared import Pt
    for t, bold, code in inline_runs(text):
        r = par.add_run(t)
        if bold:
            r.bold = True
        if code:
            r.font.name = "Consolas"
        if size:
            r.font.size = Pt(size)


def cmd_word(a):
    from docx import Document
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Cm, Pt, RGBColor

    src = need(a.file)
    blocks = parse_md(source_text(src, a))
    ea = a.font or FONT_EA
    if a.template:
        doc = Document(need(a.template))
        body = doc.element.body
        for el in list(body):
            if el.tag != qn("w:sectPr"):
                body.remove(el)
    else:
        doc = Document()
        for s in doc.sections:
            s.top_margin = s.bottom_margin = Cm(2.54)
            s.left_margin = s.right_margin = Cm(2.54)
        st = doc.styles["Normal"]
        docx_font(st.element, FONT_LATIN, ea)
        st.font.size = Pt(a.size)
        st.paragraph_format.line_spacing = 1.25
        st.paragraph_format.space_after = Pt(6)
        for name, size in (("Title", 20), ("Heading 1", 16), ("Heading 2", 14), ("Heading 3", 13), ("Heading 4", 12)):
            hs = doc.styles[name]
            docx_font(hs.element, FONT_LATIN, ea)
            hs.font.size = Pt(size)
            hs.font.bold = True
            hs.font.color.rgb = RGBColor(0, 0, 0)

    styles = {s.name for s in doc.styles}
    # 只有開頭一個 # → 當文件標題，## 以下當第 1 層標題
    h1 = [b for b in blocks if b["k"] == "h" and b["level"] == 1]
    title_mode = len(h1) == 1 and blocks and blocks[0] is h1[0]
    shift = 1 if title_mode else 0

    for b in blocks:
        k = b["k"]
        if k == "h":
            if title_mode and b is h1[0]:
                p = doc.add_paragraph(style="Title" if "Title" in styles else None)
                add_runs(p, b["text"])
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            else:
                doc.add_heading(plain(b["text"]), level=max(1, min(b["level"] - shift, 4)))
        elif k == "p":
            add_runs(doc.add_paragraph(), b["text"])
        elif k == "li":
            if b["ordered"]:
                p = doc.add_paragraph()
                add_runs(p, f"{b['marker']} {b['text']}" if b["marker"][-1] in ".)" else f"{b['marker']}{b['text']}")
                p.paragraph_format.left_indent = Cm(0.75 * (b["level"] + 1))
                p.paragraph_format.first_line_indent = Cm(-0.75)
            else:
                style = "List Bullet" + (f" {b['level'] + 1}" if b["level"] else "")
                if style in styles:
                    add_runs(doc.add_paragraph(style=style), b["text"])
                else:
                    p = doc.add_paragraph()
                    add_runs(p, "• " + b["text"])
                    p.paragraph_format.left_indent = Cm(0.75 * (b["level"] + 1))
                    p.paragraph_format.first_line_indent = Cm(-0.5)
        elif k == "table":
            rows = b["rows"]
            w = max(len(r) for r in rows)
            t = doc.add_table(rows=len(rows), cols=w)
            if "Table Grid" in styles:
                t.style = "Table Grid"
            t.alignment = WD_TABLE_ALIGNMENT.CENTER
            for ri, r in enumerate(rows):
                for ci in range(w):
                    cell = t.cell(ri, ci)
                    text = r[ci] if ci < len(r) else ""
                    p = cell.paragraphs[0]
                    add_runs(p, text)
                    if ri == 0:
                        for run in p.runs:
                            run.bold = True
                        shd = OxmlElement("w:shd")
                        shd.set(qn("w:val"), "clear")
                        shd.set(qn("w:fill"), "D9E2F3")
                        cell._tc.get_or_add_tcPr().append(shd)
                    elif to_number(plain(text)) is not None:
                        p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            doc.add_paragraph()
        elif k == "quote":
            p = doc.add_paragraph()
            add_runs(p, b["text"])
            p.paragraph_format.left_indent = Cm(1)
            for r in p.runs:
                r.italic = True
        elif k == "code":
            p = doc.add_paragraph()
            r = p.add_run(b["text"])
            r.font.name = "Consolas"
            r.font.size = Pt(10)
        elif k == "img":
            img = resolve_img(b["path"], src.parent)
            if img:
                doc.add_picture(str(img), width=Cm(15))
                doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
                if b["alt"]:
                    cap = doc.add_paragraph(b["alt"])
                    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
            else:
                doc.add_paragraph(f"［圖片：{b['alt'] or b['path']}（找不到檔案）］")
        elif k == "hr":
            doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    out = out_path(src, ".docx", a.out, a.force)
    doc.save(out)
    done(out)


# =============================================================================
#  Markdown → PowerPoint
# =============================================================================
def split_slides(blocks):
    """# = 封面（第一個）或章節頁；## = 新投影片；--- = 換頁沿用標題"""
    slides, cur, cover = [], None, None
    for b in blocks:
        if b["k"] == "h" and b["level"] == 1:
            if cover is None and not slides and cur is None:
                cover = {"kind": "cover", "title": b["text"], "sub": [], "notes": []}
            else:
                cur = None
                slides.append({"kind": "section", "title": b["text"], "items": [], "notes": []})
        elif b["k"] == "h" and b["level"] == 2:
            cur = {"kind": "content", "title": b["text"], "items": [], "notes": []}
            slides.append(cur)
        elif b["k"] == "hr":
            if cur is not None:
                cur = {"kind": "content", "title": cur["title"] + "（續）", "items": [], "notes": []}
                slides.append(cur)
        elif b["k"] == "quote":
            target = cur or (slides[-1] if slides else cover)
            if target is not None:
                target["notes"].append(re.sub(r"^(備註|講稿|notes?)[:：]\s*", "", b["text"], flags=re.I))
        elif cur is None:
            if cover is not None and not slides and b["k"] == "p":
                cover["sub"].append(b["text"])
            elif b["k"] in ("p", "li", "table", "img", "code"):
                cur = {"kind": "content", "title": "", "items": [b], "notes": []}
                slides.append(cur)
            elif b["k"] == "h":
                cur = {"kind": "content", "title": b["text"], "items": [], "notes": []}
                slides.append(cur)
        else:
            if b["k"] == "h":    # ### 小標題當粗體文字
                b = {"k": "p", "text": f"**{plain(b['text'])}**"}
            cur["items"].append(b)
    # 文字太多就拆成多張（小模型常一頁塞太多）
    out = []
    for s in slides:
        if s["kind"] != "content":
            out.append(s)
            continue
        text_items = [b for b in s["items"] if b["k"] in ("p", "li", "code")]
        others = [b for b in s["items"] if b["k"] not in ("p", "li", "code")]
        chunks, chunk, size = [], [], 0
        for b in text_items:
            n = len(plain(b["text"]))
            if chunk and (len(chunk) >= 7 or size + n > 320):
                chunks.append(chunk)
                chunk, size = [], 0
            chunk.append(b)
            size += n
        chunks.append(chunk)
        for ci, c in enumerate(chunks):
            out.append({"kind": "content", "title": s["title"] + ("（續）" if ci else ""),
                        "items": c + (others if ci == len(chunks) - 1 else []),
                        "notes": s["notes"] if ci == 0 else []})
    return cover, out


def cmd_slides(a):
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
    from pptx.util import Emu, Inches, Pt

    NAVY, ACCENT, GREY, WHITE, LIGHT = (RGBColor(0x1F, 0x3A, 0x5F), RGBColor(0x2E, 0x86, 0xAB),
                                        RGBColor(0x55, 0x5E, 0x6B), RGBColor(0xFF, 0xFF, 0xFF),
                                        RGBColor(0xEE, 0xF2, 0xF6))
    src = need(a.file)
    cover, slides = split_slides(parse_md(source_text(src, a)))
    if cover is None and not slides:
        raise UserError("Markdown 裡沒有內容：請用 # 寫封面標題、## 寫每一張投影片的標題")
    prs = Presentation(need(a.template)) if a.template else Presentation()
    if not a.template:
        prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    W, H = prs.slide_width, prs.slide_height
    blank = prs.slide_layouts[6] if len(prs.slide_layouts) > 6 else prs.slide_layouts[-1]

    def font(run, size, bold=False, color=None):
        run.font.name = PPT_FONT
        run.font.size = Pt(size)
        run.font.bold = bold
        if color is not None:
            run.font.color.rgb = color
        rpr = run._r.get_or_add_rPr()
        for tag in ("a:ea", "a:latin"):
            el = rpr.find("{http://schemas.openxmlformats.org/drawingml/2006/main}" + tag.split(":")[1])
            if el is None:
                el = rpr.makeelement("{http://schemas.openxmlformats.org/drawingml/2006/main}" + tag.split(":")[1], {})
                rpr.append(el)
            el.set("typeface", PPT_FONT)

    def text_runs(par, text, size, color=None, bold=False):
        for t, b, code in inline_runs(text):
            r = par.add_run()
            r.text = t
            font(r, size, bold or b, color)

    def box(slide, x, y, w, h):
        tb = slide.shapes.add_textbox(x, y, w, h)
        tf = tb.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_right = Inches(0.1)
        return tf

    def rect(slide, x, y, w, h, color):
        r = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, y, w, h)
        r.fill.solid()
        r.fill.fore_color.rgb = color
        r.line.fill.background()
        return r

    def notes(slide, lines):
        if lines:
            slide.notes_slide.notes_text_frame.text = "\n".join(lines)

    total = len(slides) + (1 if cover else 0)
    if cover:
        s = prs.slides.add_slide(blank)
        rect(s, 0, 0, W, H, NAVY)
        rect(s, Inches(0.8), Inches(4.1), Inches(1.5), Inches(0.08), ACCENT)
        tf = box(s, Inches(0.8), Inches(1.8), W - Inches(1.6), Inches(2.2))
        tf.vertical_anchor = MSO_ANCHOR.BOTTOM
        text_runs(tf.paragraphs[0], cover["title"], 40, WHITE, True)
        if cover["sub"]:
            tf = box(s, Inches(0.8), Inches(4.4), W - Inches(1.6), Inches(2))
            for i, line in enumerate(cover["sub"]):
                p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
                text_runs(p, line, 20, RGBColor(0xD0, 0xDC, 0xE8))
        notes(s, cover["notes"])

    for n, sd in enumerate(slides, 2 if cover else 1):
        s = prs.slides.add_slide(blank)
        if sd["kind"] == "section":
            rect(s, 0, 0, W, H, ACCENT)
            tf = box(s, Inches(0.8), Inches(2.6), W - Inches(1.6), Inches(2))
            tf.vertical_anchor = MSO_ANCHOR.MIDDLE
            text_runs(tf.paragraphs[0], sd["title"], 36, WHITE, True)
            notes(s, sd["notes"])
            continue
        # 標題列
        rect(s, 0, 0, W, Inches(1.1), NAVY)
        tf = box(s, Inches(0.6), Inches(0.15), W - Inches(1.2), Inches(0.85))
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        text_runs(tf.paragraphs[0], sd["title"], 28, WHITE, True)
        # 頁碼
        tf = box(s, W - Inches(1.5), H - Inches(0.5), Inches(1.2), Inches(0.4))
        tf.paragraphs[0].alignment = PP_ALIGN.RIGHT
        text_runs(tf.paragraphs[0], f"{n} / {total}", 12, GREY)

        texts = [b for b in sd["items"] if b["k"] in ("p", "li", "code")]
        tables = [b for b in sd["items"] if b["k"] == "table"]
        imgs = [b for b in sd["items"] if b["k"] == "img"]
        img_paths = [p for p in (resolve_img(b["path"], src.parent) for b in imgs) if p]
        top, left, bottom = Inches(1.45), Inches(0.7), H - Inches(0.7)
        width = W - Inches(1.4)
        text_w = width // 2 - Inches(0.2) if img_paths else width
        chars = sum(len(plain(b["text"])) for b in texts)
        size = 24 if len(texts) <= 4 and chars < 120 else 20 if chars < 240 else 18
        text_h = (bottom - top) if not tables else min(Inches(0.55) * len(texts) + Inches(0.2), (bottom - top) // 2)
        if texts:
            tf = box(s, left, top, text_w, text_h)
            for i, b in enumerate(texts):
                p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
                p.space_after = Pt(10)
                if b["k"] == "li":
                    lv = b["level"]
                    mark = (b["marker"] + " ") if b["ordered"] else ("• " if lv == 0 else "– ")
                    p.level = min(lv, 4)
                    sz = size - 2 * min(lv, 2)
                    r = p.add_run()
                    font(r, sz, True, ACCENT if lv == 0 else GREY)
                    r.text = mark
                    text_runs(p, b["text"], sz)
                elif b["k"] == "code":
                    r = p.add_run()
                    font(r, size - 4)
                    r.font.name = "Consolas"
                    r.text = b["text"]
                else:
                    text_runs(p, b["text"], size)
        if tables:
            t_top = top + (text_h + Inches(0.2) if texts else 0)
            avail = bottom - t_top
            for tb in tables:
                rows = tb["rows"]
                cols = max(len(r) for r in rows)
                rh = min(Inches(0.5), avail // max(len(rows), 1))
                shape = s.shapes.add_table(len(rows), cols, left, t_top, text_w, rh * len(rows))
                fsz = 16 if len(rows) <= 6 else 14 if len(rows) <= 10 else 11
                for ri, r in enumerate(rows):
                    for ci in range(cols):
                        cell = shape.table.cell(ri, ci)
                        cell.text = ""
                        par = cell.text_frame.paragraphs[0]
                        text_runs(par, r[ci] if ci < len(r) else "", fsz, WHITE if ri == 0 else None, ri == 0)
                        if ri > 0 and to_number(plain(r[ci] if ci < len(r) else "")) is not None:
                            par.alignment = PP_ALIGN.RIGHT
                        cell.fill.solid()
                        cell.fill.fore_color.rgb = NAVY if ri == 0 else (LIGHT if ri % 2 == 0 else WHITE)
                t_top += rh * len(rows) + Inches(0.2)
                avail = bottom - t_top
        if img_paths:
            x = left + (text_w + Inches(0.4) if texts or tables else 0)
            iw = (W - Inches(0.7) - x)
            ih = (bottom - top) // len(img_paths)
            for k, ip in enumerate(img_paths):
                pic = s.shapes.add_picture(str(ip), x, top + ih * k)
                scale = min(iw / pic.width, (ih - Inches(0.1)) / pic.height)
                pic.width, pic.height = int(pic.width * scale), int(pic.height * scale)
                pic.left = int(x + (iw - pic.width) / 2)
        missing = [b for b in imgs if not resolve_img(b["path"], src.parent)]
        if missing:
            sd["notes"].append("找不到圖片：" + "、".join(b["path"] for b in missing))
        notes(s, sd["notes"])

    out = out_path(src, ".pptx", a.out, a.force)
    prs.save(out)
    done(out)
    print(f"共 {total} 張投影片")


# =============================================================================
#  Markdown 表格 / CSV → Excel
# =============================================================================
def cmd_excel(a):
    from datetime import datetime
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    src = need(a.file)
    text = source_text(src, a)
    sheets = []
    if src.suffix.lower() in (".csv", ".tsv"):
        sheets.append((src.stem, list(csv.reader(io.StringIO(text), delimiter="\t" if src.suffix.lower() == ".tsv" else ","))))
    else:
        last_h = None
        for b in parse_md(text):
            if b["k"] == "h":
                last_h = plain(b["text"])
            elif b["k"] == "table":
                sheets.append((last_h or f"表{len(sheets) + 1}", [[plain(c) for c in r] for r in b["rows"]]))
                last_h = None
    if not sheets:
        raise UserError("找不到表格：Markdown 請用 | 欄 | 欄 | 的表格語法，或改給 .csv 檔")

    wb = Workbook()
    wb.remove(wb.active)
    thin = Side(style="thin", color="A6A6A6")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    used = set()
    skip_total = re.compile(r"編號|序號|項次|號碼|電話|日期|年|月|學號|代碼|No\.?|ID", re.I)
    for name, rows in sheets:
        rows = [r for r in rows if any(str(c).strip() for c in r)]
        if not rows:
            continue
        title = re.sub(r"[\[\]:*?/\\]", "", name)[:31] or "Sheet"
        base, k = title, 2
        while title in used:
            title = f"{base[:28]}_{k}"
            k += 1
        used.add(title)
        ws = wb.create_sheet(title)
        w = max(len(r) for r in rows)
        numeric_cols = set(range(w))
        for ri, r in enumerate(rows, 1):
            for ci in range(w):
                raw = str(r[ci]).strip() if ci < len(r) else ""
                cell = ws.cell(ri, ci + 1)
                cell.border = border
                if ri == 1:
                    cell.value = raw
                    cell.font = Font(bold=True, color="FFFFFF")
                    cell.fill = PatternFill("solid", fgColor="1F3A5F")
                    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
                    continue
                num = to_number(raw)
                dm = re.fullmatch(r"(\d{4})[-/.](\d{1,2})[-/.](\d{1,2})", raw)
                if num is not None:
                    cell.value = num[0]
                    cell.number_format = "0.0%" if num[1] == "%" else ("#,##0" if isinstance(num[0], int) else "#,##0.00")
                elif dm:
                    cell.value = datetime(int(dm.group(1)), int(dm.group(2)), int(dm.group(3)))
                    cell.number_format = "yyyy-mm-dd"
                    numeric_cols.discard(ci)
                else:
                    cell.value = raw
                    if raw:
                        numeric_cols.discard(ci)
                cell.alignment = Alignment(vertical="center", wrap_text=True)
        if a.total and len(rows) > 2:
            tr = len(rows) + 1
            label_col = next((c for c in range(w) if c not in numeric_cols), 0)
            ws.cell(tr, label_col + 1, "合計").font = Font(bold=True)
            for ci in range(w):
                head = str(rows[0][ci]) if ci < len(rows[0]) else ""
                cell = ws.cell(tr, ci + 1)
                cell.border = border
                cell.fill = PatternFill("solid", fgColor="EEF2F6")
                if ci in numeric_cols and ci != label_col and not skip_total.search(head):
                    L = get_column_letter(ci + 1)
                    cell.value = f"=SUM({L}2:{L}{tr - 1})"
                    cell.number_format = ws.cell(2, ci + 1).number_format or "#,##0"
                    cell.font = Font(bold=True)
        for ci in range(w):
            longest = max((sum(2 if is_cjk(ch) else 1 for ch in str(r[ci])) for r in rows if ci < len(r)), default=4)
            ws.column_dimensions[get_column_letter(ci + 1)].width = min(max(longest + 2, 8), 60)
        ws.freeze_panes = "A2"
    out = out_path(src, ".xlsx", a.out, a.force)
    wb.save(out)
    done(out)
    print(f"共 {len(wb.worksheets)} 個工作表：{'、'.join(ws.title for ws in wb.worksheets)}")


# =============================================================================
#  PDF
# =============================================================================
def cmd_pdf2word(a):
    import pymupdf as fitz
    src = need(a.file)
    with fitz.open(src) as d:
        n = d.page_count
        pages = page_list(a.pages, n)
        scanned = all(len(d[i].get_text("text").strip()) < 20 for i in pages)
    if scanned:
        print("【注意】這份 PDF 讀不到文字（可能是掃描檔），轉出來的 Word 只會是圖片，無法編輯文字。")
    import logging
    logging.disable(logging.WARNING)
    from pdf2docx import Converter
    out = out_path(src, ".docx", a.out, a.force)
    print(f"轉換中（共 {len(pages)} 頁，頁數多會比較久）…", flush=True)
    cv = Converter(str(src))
    try:
        cv.convert(str(out), pages=pages)
    finally:
        cv.close()
    done(out)


def cmd_pdf(a):
    import pymupdf as fitz
    op = a.op
    if op == "merge":
        files = [need(f) for f in a.files]
        out = out_path(files[0], ".pdf", a.out, a.force, "_合併")
        res = fitz.open()
        for f in files:
            if f.suffix.lower() == ".pdf":
                with fitz.open(f) as d:
                    res.insert_pdf(d)
            else:
                with fitz.open(f) as img:
                    res.insert_pdf(fitz.open("pdf", img.convert_to_pdf()))
        res.save(out, garbage=3, deflate=True)
        done(out)
        print(f"合併 {len(files)} 個檔案，共 {res.page_count} 頁")
        return
    if op == "fromimg":
        files = [need(f) for f in a.files]
        out = out_path(files[0], ".pdf", a.out, a.force)
        res = fitz.open()
        for f in files:
            with fitz.open(f) as img:
                rect = img[0].rect
                pdf = fitz.open("pdf", img.convert_to_pdf())
            # 縮放到 A4 寬度內，方便列印
            a4 = fitz.paper_rect("a4")
            page = res.new_page(width=a4.width, height=a4.height)
            scale = min((a4.width - 40) / rect.width, (a4.height - 40) / rect.height, 1)
            w, h = rect.width * scale, rect.height * scale
            page.show_pdf_page(fitz.Rect((a4.width - w) / 2, (a4.height - h) / 2, (a4.width + w) / 2, (a4.height + h) / 2), pdf, 0)
        res.save(out, garbage=3, deflate=True)
        done(out)
        print(f"共 {res.page_count} 頁（每張圖片一頁，置中縮放到 A4）")
        return

    src = need(a.file)
    doc = fitz.open(src)
    n = doc.page_count
    if op == "pages":
        pages = page_list(a.range, n)
        doc.select(pages)
        out = out_path(src, ".pdf", a.out, a.force, f"_第{a.range}頁".replace(",", "_"))
    elif op == "delete":
        drop = set(page_list(a.range, n))
        doc.select([i for i in range(n) if i not in drop])
        out = out_path(src, ".pdf", a.out, a.force, "_刪頁")
    elif op == "rotate":
        for i in page_list(a.pages, n):
            doc[i].set_rotation((doc[i].rotation + a.degrees) % 360)
        out = out_path(src, ".pdf", a.out, a.force, "_旋轉")
    elif op == "topng":
        pages = page_list(a.pages, n)
        folder = Path(a.out) if a.out else src.with_name(src.stem + "_圖片")
        folder.mkdir(parents=True, exist_ok=True)
        for i in pages:
            doc[i].get_pixmap(dpi=a.dpi).save(folder / f"{src.stem}_p{i + 1:03d}.png")
        print(f"已產生 {len(pages)} 張圖片：{folder.resolve()}")
        return
    elif op == "compress":
        before = src.stat().st_size
        try:
            doc.rewrite_images(dpi_threshold=a.dpi + 10, dpi_target=a.dpi, quality=70)
        except Exception:
            pass   # 舊版 PyMuPDF 沒有 rewrite_images，只做結構壓縮
        out = out_path(src, ".pdf", a.out, a.force, "_壓縮")
        doc.save(out, garbage=4, deflate=True, clean=True)
        after = Path(out).stat().st_size
        done(out)
        print(f"{before / 1e6:.2f} MB → {after / 1e6:.2f} MB")
        return
    doc.save(out, garbage=3, deflate=True)
    done(out)
    print(f"共 {doc.page_count} 頁")


# =============================================================================
#  Office → PDF
# =============================================================================
def cmd_topdf(a):
    src = need(a.file).resolve()
    ext = src.suffix.lower()
    out = out_path(src, ".pdf", a.out, a.force).resolve()
    soffice = shutil.which("soffice") or shutil.which("libreoffice")
    if not soffice and os.name == "nt":
        for base in (os.environ.get("ProgramFiles", ""), os.environ.get("ProgramFiles(x86)", "")):
            cand = Path(base) / "LibreOffice" / "program" / "soffice.exe"
            if cand.exists():
                soffice = str(cand)
    if os.name == "nt":
        try:
            office_to_pdf(src, out, ext)
            done(out)
            return
        except UserError:
            if not soffice:
                raise
    if not soffice:
        raise UserError("轉 PDF 需要 LibreOffice：sudo apt install libreoffice（Windows 有裝 Office 也可以）")
    tmp = out.parent / ".labdoc-tmp"
    tmp.mkdir(exist_ok=True)
    try:
        subprocess.run([soffice, "--headless", "--convert-to", "pdf", "--outdir", str(tmp), str(src)],
                       check=True, capture_output=True, timeout=300)
        shutil.move(str(tmp / (src.stem + ".pdf")), str(out))
    finally:
        shutil.rmtree(tmp, ignore_errors=True)
    done(out)


# 用 PowerShell 操作已安裝的 Microsoft Office（COM），不必另外裝 pywin32
OFFICE_PS = {
    "word": "$a = New-Object -ComObject Word.Application; $a.Visible = $false; $a.DisplayAlerts = 0; "
            "try { $d = $a.Documents.Open($env:LABDOC_SRC, $false, $true); $d.ExportAsFixedFormat($env:LABDOC_OUT, 17); "
            "$d.Close($false) } finally { $a.Quit() }",
    "powerpoint": "$a = New-Object -ComObject PowerPoint.Application; "
                  "try { $d = $a.Presentations.Open($env:LABDOC_SRC, -1, 0, 0); $d.SaveAs($env:LABDOC_OUT, 32); "
                  "$d.Close() } finally { $a.Quit() }",
    "excel": "$a = New-Object -ComObject Excel.Application; $a.Visible = $false; $a.DisplayAlerts = $false; "
             "try { $d = $a.Workbooks.Open($env:LABDOC_SRC, 0, $true); $d.ExportAsFixedFormat(0, $env:LABDOC_OUT); "
             "$d.Close($false) } finally { $a.Quit() }",
}


def office_to_pdf(src, out, ext):
    app = ("word" if ext in (".doc", ".docx", ".rtf", ".odt", ".txt") else
           "powerpoint" if ext in (".ppt", ".pptx") else
           "excel" if ext in (".xls", ".xlsx", ".xlsm") else None)
    if not app:
        raise UserError(f"不支援 {ext} 轉 PDF")
    # 路徑用環境變數傳，避免中文或空白在指令列被拆開
    env = dict(os.environ, LABDOC_SRC=str(src), LABDOC_OUT=str(out))
    r = subprocess.run(["powershell", "-NoProfile", "-NonInteractive", "-Command", OFFICE_PS[app]],
                       env=env, capture_output=True, timeout=300)
    if r.returncode != 0 or not Path(out).exists():
        raise UserError("轉 PDF 需要 Microsoft Office 或 LibreOffice（這台的 Office 無法使用）")


# =============================================================================
#  tw：簡體字、大陸用語、半形標點 → 台灣繁體
#  用 OpenCC（s2twp）查字典，比小模型自己判斷可靠；分三類：
#    簡體字   一定要改（--fix 會改）
#    大陸用語 建議改（--fix 會改）
#    看上下文 台灣也常用、只在某些意思下才算大陸用語（只提醒，不自動改）
# =============================================================================
# OpenCC 沒收錄的常見大陸用語，或 OpenCC 的建議不是台灣常用說法
TW_EXTRA = {
    "質量": "品質", "響應": "回應", "反饋": "回饋", "公交車": "公車", "公交": "公車", "數字化": "數位化",
    "帖子": "貼文", "靠譜": "可靠", "套餐": "方案", "營銷": "行銷", "策劃": "企劃", "運營": "營運",
    "賬號": "帳號", "賬戶": "帳戶", "渠道": "管道", "調研": "調查", "對接": "接洽", "移動端": "行動裝置",
    "攝像頭": "鏡頭", "觸屏": "觸控螢幕", "充電寶": "行動電源", "外賣": "外送", "地鐵": "捷運",
    "點贊": "按讚", "二維碼": "QR Code", "本科生": "大學部學生", "簡歷": "履歷", "入職": "到職",
    "報銷": "核銷", "打車": "搭計程車", "酒店": "飯店", "衛生間": "洗手間", "方便麵": "泡麵",
    "土豆": "馬鈴薯", "西紅柿": "番茄", "冰激凌": "冰淇淋", "摩托車": "機車", "集裝箱": "貨櫃",
    "激光": "雷射", "筆記本電腦": "筆電", "視頻會議": "視訊會議", "在線": "線上", "屏蔽": "封鎖",
    "博主": "部落客", "網民": "網友", "給力": "很棒", "搞定": "完成", "立馬": "馬上", "牛逼": "厲害",
}
# 台灣也常用的詞：只有在特定意思下才是大陸用語 → 只提醒
TW_CONTEXT = {
    "支持": "指技術或軟體支援時用「支援」；表示贊成、鼓勵時「支持」正確",
    "項目": "指計畫時用「專案」；指條列項目、評分項目時「項目」正確",
    "程序": "指電腦程式時用「程式」；行政、法律程序時「程序」正確",
    "窗口": "指電腦視窗時用「視窗」；服務窗口、聯絡窗口時「窗口」正確",
    "數據": "台灣多用「資料」，統計數字也可用「數據」",
    "優化": "台灣也用「優化」，正式文件可寫「最佳化」或「改善」",
    "領導": "指人時台灣用「長官、主管」；動詞「領導」正確",
    "水平": "指程度時台灣用「水準」；水平線、水平方向時正確",
    "落地": "指計畫實行時台灣用「落實、實施」",
    "賦能": "大陸流行語，台灣可寫「協助、培力」",
    "閉環": "大陸流行語，台灣可寫「完整流程、回饋循環」",
    "抓手": "大陸流行語，台灣可寫「著力點、切入點」",
}
_PUNCT = {",": "，", ";": "；", ":": "：", "?": "？", "!": "！"}
_CJK = r"[\u3400-\u9fff\uf900-\ufaff]"
# 中文字後面的半形標點（後面接中文或行尾）；只吃同一行的空白，不能跨行
_PUNCT_RE = re.compile(rf"(?<={_CJK})[ \t]*([,;:?!])[ \t]*(?={_CJK}|$)", re.M)
# 「台」「臺」在台灣都正確，OpenCC 會把台改成臺 → 不動它
_TW_KEEP_CHARS = "台"


class TwChecker:
    def __init__(self):
        try:
            import opencc
            from opencc import OpenCC
        except ModuleNotFoundError:
            raise UserError("缺少 OpenCC，請重新執行 install-client 安裝文件工具")
        self.chars = OpenCC("s2tw")     # 簡體字 → 繁體字（含台灣異體字）
        # 大陸用語：OpenCC 的台灣詞典（繁體寫法的大陸詞 → 台灣詞）＋ 補充詞表，扣掉要看上下文的詞
        terms = {}
        dict_path = Path(opencc.__file__).parent / "dictionary" / "TWPhrases.txt"
        for line in dict_path.read_text(encoding="utf-8").splitlines():
            parts = line.split("\t")
            if len(parts) == 2 and parts[1].split():
                terms[parts[0]] = parts[1].split()[0]
        terms.update(TW_EXTRA)
        for w in TW_CONTEXT:
            terms.pop(w, None)
        self.terms = {k: v for k, v in terms.items() if k != v}
        self.maxlen = max(map(len, self.terms))
        self.ctx_re = re.compile("|".join(map(re.escape, sorted(TW_CONTEXT, key=len, reverse=True))))

    def to_trad(self, s):
        parts = re.split(f"([{_TW_KEEP_CHARS}])", s)
        return "".join(x if x and x in _TW_KEEP_CHARS else self.chars.convert(x) for x in parts)

    def match_terms(self, s):
        """由左到右最長比對，回傳 [(開始, 結束, 台灣用語)]；「看上下文」的詞整個跳過，不拆開比對"""
        out, i, n = [], 0, len(s)
        while i < n:
            for L in range(min(self.maxlen, n - i), 0, -1):
                w = s[i:i + L]
                if w in TW_CONTEXT:
                    i += L
                    break
                if w in self.terms:
                    out.append((i, i + L, self.terms[w]))
                    i += L
                    break
            else:
                i += 1
        return out

    def stages(self, s):
        """[(類別, 轉換前, 轉換後, [(開始, 結束, 新字)])]：簡體字 → 大陸用語 → 半形標點"""
        import difflib
        s1 = self.to_trad(s)
        ch = [(i1, i2, s1[j1:j2]) for tag, i1, i2, j1, j2 in
              difflib.SequenceMatcher(None, s, s1, autojunk=False).get_opcodes() if tag != "equal"]
        t = self.match_terms(s1)
        s2 = s1
        for a, b, w in reversed(t):
            s2 = s2[:a] + w + s2[b:]
        pu = [(m.start(), m.end(), _PUNCT[m.group(1)]) for m in _PUNCT_RE.finditer(s2)]
        s3 = _PUNCT_RE.sub(lambda m: _PUNCT[m.group(1)], s2)
        return [("簡體字", s, s1, ch), ("大陸用語", s1, s2, t), ("半形標點", s2, s3, pu)]

    def fix(self, s):
        return self.stages(s)[-1][2]

    def findings(self, s):
        out = []
        for kind, a, _, spans in self.stages(s):
            for i1, i2, new in spans:
                src = a[i1:i2].strip(" \t") if kind == "半形標點" else a[i1:i2]
                out.append((kind, src, new, a[max(0, i1 - 8):i2 + 8].replace("\n", " ")))
        s1 = self.to_trad(s)
        covered = [(i1, i2) for i1, i2, _ in self.match_terms(s1)]   # 已經在大陸用語裡改掉的詞（例：數據庫）
        for m in self.ctx_re.finditer(s1):
            if any(i1 <= m.start() and m.end() <= i2 for i1, i2 in covered):
                continue
            out.append(("看上下文", m.group(0), TW_CONTEXT[m.group(0)], s1[max(0, m.start() - 8):m.end() + 8].replace("\n", " ")))
        return out


def tw_units(p):
    """把檔案拆成 (位置, 文字)：Word 段落、PPT 每張、PDF 每頁、文字檔每行"""
    ext = p.suffix.lower()
    if ext == ".docx":
        from docx import Document
        doc = Document(p)
        n = 0
        for para in doc.paragraphs:
            n += 1
            if para.text.strip():
                yield f"第 {n} 段", para.text
        for ti, t in enumerate(doc.tables, 1):
            for ri, r in enumerate(t.rows, 1):
                for ci, c in enumerate(r.cells, 1):
                    if c.text.strip():
                        yield f"表{ti} 第{ri}列第{ci}欄", c.text
    elif ext == ".pptx":
        from pptx import Presentation
        for n, slide in enumerate(Presentation(p).slides, 1):
            for sh in slide.shapes:
                if sh.has_text_frame and sh.text_frame.text.strip():
                    yield f"第 {n} 張", sh.text_frame.text
                if getattr(sh, "has_table", False) and sh.has_table:
                    for r in sh.table.rows:
                        for c in r.cells:
                            if c.text.strip():
                                yield f"第 {n} 張（表格）", c.text
    elif ext == ".pdf":
        import pymupdf as fitz
        with fitz.open(p) as d:
            for i in range(d.page_count):
                yield f"第 {i + 1} 頁", d[i].get_text("text")
    elif ext in (".xlsx", ".xlsm"):
        from openpyxl import load_workbook
        for ws in load_workbook(p, read_only=True).worksheets:
            for row in ws.iter_rows():
                for c in row:
                    if isinstance(c.value, str) and c.value.strip():
                        yield f"{ws.title}!{c.coordinate}", c.value
    else:
        for n, line in enumerate(read_text_file(p).splitlines(), 1):
            if line.strip():
                yield f"第 {n} 行", line


def tw_fix_file(p, ck, out):
    """另存修正版：Word / PPT 逐段改字，保留原本格式"""
    ext = p.suffix.lower()

    def fix_runs(runs):
        # 先試整段一起改（詞可能被拆在不同 run），改完長度不變就照位置放回；否則逐個 run 改
        runs = list(runs)
        whole = "".join(r.text for r in runs)
        fixed = ck.fix(whole)
        if fixed == whole:
            return
        if len(fixed) == len(whole):
            pos = 0
            for r in runs:
                r.text, pos = fixed[pos:pos + len(r.text)], pos + len(r.text)
        else:
            for r in runs:
                r.text = ck.fix(r.text)

    if ext == ".docx":
        from docx import Document
        doc = Document(p)
        paras = list(doc.paragraphs)
        for t in doc.tables:
            for r in t.rows:
                for c in r.cells:
                    paras.extend(c.paragraphs)
        for s in doc.sections:
            for part in (s.header, s.footer):
                paras.extend(part.paragraphs)
        for para in paras:
            fix_runs(para.runs)
        doc.save(out)
    elif ext == ".pptx":
        from pptx import Presentation
        prs = Presentation(p)
        for slide in prs.slides:
            frames = [sh.text_frame for sh in slide.shapes if sh.has_text_frame]
            for sh in slide.shapes:
                if getattr(sh, "has_table", False) and sh.has_table:
                    frames += [c.text_frame for r in sh.table.rows for c in r.cells]
            if slide.has_notes_slide:
                frames.append(slide.notes_slide.notes_text_frame)
            for tf in frames:
                for para in tf.paragraphs:
                    fix_runs(para.runs)
        prs.save(out)
    elif ext in (".pdf", ".xlsx", ".xlsm"):
        raise UserError(f"{ext} 不能直接修正：只列出問題；PDF 可先用 labdoc pdf2word 轉成 Word 再修正")
    else:
        Path(out).write_text(ck.fix(read_text_file(p)), encoding="utf-8")


def source_text(src, a):
    """word / slides / excel 讀 Markdown：先把簡體字、大陸用語、半形標點改成台灣用法（--keep-terms 不改）"""
    text = read_text_file(src)
    if getattr(a, "keep_terms", False):
        return text
    try:
        ck = TwChecker()
    except UserError:
        return text
    fixed = ck.fix(text)
    if fixed != text:
        from collections import Counter
        c = Counter((src_, dst) for kind, src_, dst, _ in ck.findings(text) if kind in ("簡體字", "大陸用語"))
        if c:
            print("已改成台灣用語：" + "、".join(f"{x}→{y}" for (x, y), _ in c.most_common(15))
                  + (" …" if len(c) > 15 else ""))
    return fixed


def cmd_tw(a):
    ck = TwChecker()
    p = need(a.file)
    rows = []
    for where, text in tw_units(p):
        for kind, src, dst, ctx in ck.findings(text):
            rows.append((kind, where, src, dst, ctx))
    order = {"簡體字": 0, "大陸用語": 1, "半形標點": 2, "看上下文": 3}
    rows.sort(key=lambda r: order[r[0]])
    if not rows:
        print("沒有發現簡體字、大陸用語或半形標點。")
    else:
        from collections import Counter
        counts = Counter(r[0] for r in rows)
        print("檢查結果：" + "、".join(f"{k} {v} 處" for k, v in sorted(counts.items(), key=lambda kv: order[kv[0]])))
        terms = Counter((r[0], r[2], r[3]) for r in rows if r[0] in ("簡體字", "大陸用語"))
        if terms:
            print("\n要改的字詞（次數）：" + "、".join(f"{s}→{d}×{n}" if n > 1 else f"{s}→{d}"
                                              for (_, s, d), n in terms.most_common()))
        shown = rows[:a.max]
        print("\n| 類別 | 位置 | 原文 | 建議 | 前後文 |\n|---|---|---|---|---|")
        for kind, where, src, dst, ctx in shown:
            print(f"| {kind} | {where} | {src} | {dst} | …{ctx.replace('|', '／').strip()}… |")
        if len(rows) > len(shown):
            print(f"\n【只列出前 {len(shown)} 處，共 {len(rows)} 處；可用 --max 調整】")
        print("\n說明：簡體字、大陸用語、半形標點用 --fix 會自動修正；「看上下文」請依文意自行判斷，--fix 不會改。")
    if a.fix and rows:
        out = out_path(p, p.suffix.lower(), a.out, a.force, "_台灣用語")
        tw_fix_file(p, ck, out)
        done(out)


# =============================================================================
#  指令列
# =============================================================================
def main(argv=None):
    ap = argparse.ArgumentParser(prog="labdoc", description=__doc__.split("\n\n")[1],
                                 formatter_class=argparse.RawDescriptionHelpFormatter)
    sub = ap.add_subparsers(dest="cmd", metavar="指令")

    def common(p, out=True):
        if out:
            p.add_argument("-o", "--out", help="輸出檔名（預設放在輸入檔旁邊）")
        p.add_argument("--force", action="store_true", help="輸出檔已存在時直接覆蓋")

    p = sub.add_parser("read", help="讀 PDF / Word / PPT / Excel / CSV / 文字檔，印出純文字")
    p.add_argument("file")
    p.add_argument("--pages", help="PDF 只讀這些頁，例如 1-3,5")
    p.add_argument("--max", type=int, default=MAX_CHARS, help=f"最多印出幾個字（預設 {MAX_CHARS}）")
    p.add_argument("-o", "--out", help="全文寫到這個檔案，不印出來")
    p.set_defaults(fn=cmd_read)

    p = sub.add_parser("pdf2word", help="PDF 轉 Word（保留版面）")
    p.add_argument("file")
    p.add_argument("--pages", help="只轉這些頁，例如 1-3,5")
    common(p)
    p.set_defaults(fn=cmd_pdf2word)

    p = sub.add_parser("word", help="Markdown → Word")
    p.add_argument("file")
    p.add_argument("--template", help="套用 Word 範本（沿用它的樣式、頁首頁尾）")
    p.add_argument("--font", help=f"中文字型（預設 {FONT_EA}）")
    p.add_argument("--size", type=float, default=12, help="內文字級（預設 12）")
    p.add_argument("--keep-terms", action="store_true", help="不要自動把簡體字、大陸用語改成台灣用語")
    common(p)
    p.set_defaults(fn=cmd_word)

    p = sub.add_parser("slides", help="Markdown → PowerPoint（# 封面、## 每張投影片、> 講者備註）")
    p.add_argument("file")
    p.add_argument("--template", help="PowerPoint 範本（用它的投影片大小與母片）")
    p.add_argument("--keep-terms", action="store_true", help="不要自動把簡體字、大陸用語改成台灣用語")
    common(p)
    p.set_defaults(fn=cmd_slides)

    p = sub.add_parser("excel", help="Markdown 表格 / CSV → Excel")
    p.add_argument("file")
    p.add_argument("--total", action="store_true", help="每個表格最後加「合計」列（用 SUM 公式）")
    p.add_argument("--keep-terms", action="store_true", help="不要自動把簡體字、大陸用語改成台灣用語")
    common(p)
    p.set_defaults(fn=cmd_excel)

    p = sub.add_parser("topdf", help="Word / PPT / Excel → PDF")
    p.add_argument("file")
    common(p)
    p.set_defaults(fn=cmd_topdf)

    p = sub.add_parser("pdf", help="PDF 工具：merge 合併、pages 取出頁、delete 刪頁、rotate 旋轉、"
                                   "fromimg 圖片轉 PDF、topng 轉圖片、compress 壓縮")
    ops = p.add_subparsers(dest="op", metavar="動作", required=True)
    q = ops.add_parser("merge", help="合併多個 PDF（也可混圖片），依給的順序")
    q.add_argument("files", nargs="+")
    common(q)
    q = ops.add_parser("fromimg", help="多張圖片合成一個 PDF（每張一頁，A4）")
    q.add_argument("files", nargs="+")
    common(q)
    q = ops.add_parser("pages", help="取出指定頁成新檔")
    q.add_argument("file")
    q.add_argument("range", help="例如 1-3,5")
    common(q)
    q = ops.add_parser("delete", help="刪除指定頁")
    q.add_argument("file")
    q.add_argument("range", help="例如 2,4-5")
    common(q)
    q = ops.add_parser("rotate", help="旋轉頁面")
    q.add_argument("file")
    q.add_argument("degrees", type=int, choices=[90, 180, 270, -90])
    q.add_argument("--pages", help="只轉這些頁（預設全部）")
    common(q)
    q = ops.add_parser("topng", help="每頁轉成 PNG 圖片")
    q.add_argument("file")
    q.add_argument("--pages")
    q.add_argument("--dpi", type=int, default=150)
    q.add_argument("-o", "--out", help="輸出資料夾")
    q = ops.add_parser("compress", help="壓縮（降低圖片解析度），上傳有大小限制時用")
    q.add_argument("file")
    q.add_argument("--dpi", type=int, default=120, help="圖片降到這個解析度（預設 120）")
    common(q)
    p.set_defaults(fn=cmd_pdf)

    p = sub.add_parser("tw", help="檢查簡體字、大陸用語、半形標點（台灣繁體）")
    p.add_argument("file")
    p.add_argument("--fix", action="store_true", help="另存修正版（原檔名_台灣用語），原檔不動")
    p.add_argument("--max", type=int, default=150, help="最多列出幾處（預設 150）")
    common(p)
    p.set_defaults(fn=cmd_tw)

    a = ap.parse_args(argv)
    if not getattr(a, "fn", None):
        ap.print_help()
        return 0
    try:
        a.fn(a)
    except UserError as e:
        print(f"錯誤：{e}", file=sys.stderr)
        return 1
    except ModuleNotFoundError as e:
        print(f"錯誤：缺少 Python 套件 {e.name}，請重新執行 install-client 安裝文件工具", file=sys.stderr)
        return 1
    return 0


if __name__ == "__main__":
    sys.exit(main())
